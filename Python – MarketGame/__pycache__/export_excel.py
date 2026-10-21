"""
Génère un classeur Excel qui refait, avec des formules visibles, toutes les
étapes du modèle Python et du tableau de bord du Market Game 2027.

Ce que Python prépare (colonnes sur fond bleu clair dans l'Excel) :
  - la forme de consommation 2027 à température normale (profils Enedis
    dé-météorisés, moyenne 2023-2025, alignés sur le calendrier 2027) ;
  - les températures normales et les gradients 2027 ;
  - l'écart de température de chaque année météo (2022 à 2025) ;
  - les prix day-ahead historiques de chaque année, replacés sur 2027 ;
  - les données journalières 2022-2025 pour la régression prix-volume.

Ce que l'Excel calcule lui-même, avec des formules :
  consommation et coefficient météo CM, volumes, anomalies, pente b,
  scénarios de prix spot, prix T2 anticipé (alpha), stratégies en MW,
  marges, règle du Board, décision T2, position ouverte du 01/01/2027,
  et (si config.DA_0101 existe) le résultat du day-ahead du 01/01/2027 :
  exécution de notre ordre, coût du day-ahead et marge de la journée.

Lancement : python export_excel.py
            (crée Market_Game_2027.xlsx à côté de ce script)
"""

import sys
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
    from openpyxl.chart.trendline import Trendline
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.workbook.properties import CalcProperties
except ImportError:
    raise SystemExit("Il manque le module openpyxl : lancer  pip install openpyxl")

import config
import export_dashboard as ed
from calendrier import construire_correspondance, jours_feries_fr, remapper_serie
from enedis_api import (
    charger_ou_actualiser,
    nettoyer_gradients,
    nettoyer_profil,
    nettoyer_temperatures,
    recuperer_gradients,
    recuperer_profil,
    recuperer_temperatures,
)
from load_2027 import construire_base_2027, construire_loads_scenarios
from load_model import (
    calculer_load,
    calculer_puissance_moyenne,
    construire_load_finale,
    fusionner_load_normale,
)
from prix_api import nettoyer_prix, recuperer_prix

DOSSIER = Path(__file__).resolve().parent
SORTIE = DOSSIER / "Market_Game_2027.xlsx"

NOMS_P = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Q3 (juil.-sept.)", "Q4 (oct.-déc.)"]
PROD_P = ["Jan-27", "Feb-27", "Mar-27", "Apr-27", "May-27", "Jun-27", "Q3-27", "Q4-27"]
CLES_MOIS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
PERIODES_Q = {1: [1, 2, 3], 2: [4, 5, 6], 3: [7], 4: [8]}


# ======================================================================
# 1. DONNÉES (mêmes calculs que export_dashboard.py)
# ======================================================================

def preparer_donnees():
    print("Données du tableau de bord (paramètres, stratégies)...")
    D = ed.construire_donnees()

    annees = [int(a) for a in D["meta"]["annees"]]
    annee_cible = getattr(config, "ANNEE_CIBLE", 2027)
    seuil = getattr(config, "SEUIL_CHAUFFAGE", 15.0)
    forcer = getattr(ed, "FORCER_MAJ_API", False)

    print("\nSéries horaires et journalières pour l'Excel...")
    res1 = charger_ou_actualiser("res1.csv", lambda: recuperer_profil("RES1_BASE"), nettoyer_profil, forcer)
    pro1 = charger_ou_actualiser("pro1.csv", lambda: recuperer_profil("PRO1_BASE"), nettoyer_profil, forcer)
    temperatures = charger_ou_actualiser("temperatures.csv", recuperer_temperatures, nettoyer_temperatures, forcer)
    gradients = charger_ou_actualiser("gradients.csv", recuperer_gradients, nettoyer_gradients, forcer)

    p_res1 = calculer_puissance_moyenne(config.NB_RES1, config.THETA_RES1, config.PUISSANCE_RES1)
    p_pro1 = calculer_puissance_moyenne(config.NB_PRO1, config.THETA_PRO1, config.PUISSANCE_PRO1)
    res1["load_normal_kw"] = calculer_load(res1["coefficient_dynamique"], p_res1)
    pro1["load_normal_kw"] = calculer_load(pro1["coefficient_dynamique"], p_pro1)

    base = construire_base_2027(res1, pro1, gradients, temperatures,
                                annee_ref=config.ANNEE_PROFIL_REF, annee_cible=annee_cible,
                                seuil_chauffage=seuil)
    loads = construire_loads_scenarios(base, temperatures, annees, annee_cible, seuil)

    prix = charger_ou_actualiser("prix.csv", recuperer_prix, nettoyer_prix, forcer)
    prix = prix[prix["horodate"].dt.year.isin(annees)]
    prix = prix.set_index("horodate")["prix"].resample("h").mean().reset_index()

    # Grille horaire 2027 (la même que celle du tableau de bord)
    idx = pd.date_range(f"{annee_cible}-01-01 00:00", f"{annee_cible}-12-31 23:59",
                        freq="h", tz="Europe/Paris")

    def horaire(df, col):
        s = df.set_index("horodate")[col].resample("h").mean()
        return s.reindex(idx).interpolate(limit_direction="both")

    H = pd.DataFrame(index=idx)
    H["coef_res1"] = horaire(base, "load_normal_res1_kw") / p_res1
    H["coef_pro1"] = horaire(base, "load_normal_pro1_kw") / p_pro1
    H["tn"] = horaire(base, "temperature_normale")
    H["g_res1"] = horaire(base, "gradient_res1")
    H["g_pro1"] = horaire(base, "gradient_pro1")
    for a in annees:
        H[f"dt_{a}"] = horaire(loads[a], "anomalie_temperature")
        hist = prix[prix["horodate"].dt.year == a]
        pr = remapper_serie(hist, ["prix"], construire_correspondance(a, annee_cible),
                            annee_cible, freq="h")
        H[f"prix_{a}"] = pr.set_index("horodate")["prix"].reindex(idx).interpolate(limit_direction="both")

    # Données journalières pour la régression prix-volume
    lf = construire_load_finale(fusionner_load_normale(res1, pro1), temperatures, gradients, seuil)
    lf["date"] = lf["horodate"].dt.date
    vol = lf.groupby("date")[["load_normal_total_mw", "load_final_total_mw"]].mean()
    p = prix.copy()
    p["date"] = p["horodate"].dt.date
    pj = p.groupby("date")["prix"].mean().rename("prix_jour")
    J = vol.join(pj, how="inner").reset_index()
    J = J[[d.year in annees for d in J["date"]]].reset_index(drop=True)

    feries = sorted(d for a in annees for d in jours_feries_fr(a).values())
    return D, H, J, feries, annees


# ======================================================================
# 2. OUTILS DE MISE EN FORME
# ======================================================================

POLICE = "Arial"
BLEU_FONCE = "1F3864"
F_TITRE = Font(name=POLICE, size=14, bold=True, color=BLEU_FONCE)
F_SOUS = Font(name=POLICE, size=10, italic=True, color="595959")
F_SECTION = Font(name=POLICE, size=11, bold=True, color="FFFFFF")
F_ENTETE = Font(name=POLICE, size=10, bold=True)
F_NORMAL = Font(name=POLICE, size=10)
F_INPUT = Font(name=POLICE, size=10, color="0000FF")
F_RESULTAT = Font(name=POLICE, size=10, bold=True)
FILL_SECTION = PatternFill("solid", fgColor=BLEU_FONCE)
FILL_ENTETE = PatternFill("solid", fgColor="D9E1F2")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_DONNEE = PatternFill("solid", fgColor="EAF1FB")
FILL_RESULTAT = PatternFill("solid", fgColor="E2EFDA")
BORD = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="center")

FMT = {
    "mw": "0.00", "mwh": "#,##0", "eur": "#,##0", "prix": "0.00", "pct": "0.0%",
    "coef": "0.0000", "temp": "0.0", "dt": "dd/mm/yyyy hh:mm", "date": "dd/mm/yyyy",
    "b": "0.000", "int": "0",
}


def q(s):
    return "'" + s.replace("'", "''") + "'"


def A(s, c, r):
    """Référence absolue à une cellule d'une autre feuille."""
    return f"{q(s)}!${L(c)}${r}"


def RG(s, c1, r1, c2, r2):
    """Référence absolue à une plage d'une autre feuille."""
    return f"{q(s)}!${L(c1)}${r1}:${L(c2)}${r2}"


def ecrire(ws, r, c, v, fmt=None, font=None, fill=None, bold=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font or (F_RESULTAT if bold else F_NORMAL)
    if fmt:
        cell.number_format = FMT.get(fmt, fmt)
    if fill:
        cell.fill = fill
    return cell


def entree(ws, r, c, v, fmt=None):
    """Hypothèse modifiable : police bleue sur fond jaune."""
    return ecrire(ws, r, c, v, fmt, F_INPUT, FILL_INPUT)


def donnee(ws, r, c, v, fmt=None):
    """Donnée préparée par Python : fond bleu clair."""
    return ecrire(ws, r, c, v, fmt, F_NORMAL, FILL_DONNEE)


def resultat(ws, r, c, v, fmt=None):
    return ecrire(ws, r, c, v, fmt, F_RESULTAT, FILL_RESULTAT)


def titre(ws, texte, sous=None):
    ws.cell(row=1, column=1, value=texte).font = F_TITRE
    if sous:
        ws.cell(row=2, column=1, value=sous).font = F_SOUS


def section(ws, r, texte, c1=1, c2=12):
    for c in range(c1, c2 + 1):
        ws.cell(row=r, column=c).fill = FILL_SECTION
    ws.cell(row=r, column=c1, value=texte).font = F_SECTION


def entetes(ws, r, c, libelles, hauteur=30):
    for i, t in enumerate(libelles):
        cell = ws.cell(row=r, column=c + i, value=t)
        cell.font = F_ENTETE
        cell.fill = FILL_ENTETE
        cell.alignment = WRAP
    ws.row_dimensions[r].height = hauteur


def note(ws, r, c, texte):
    ws.cell(row=r, column=c, value=texte).font = F_SOUS


def largeurs(ws, valeurs):
    for i, w in enumerate(valeurs, start=1):
        ws.column_dimensions[L(i)].width = w


def nommer(wb, nom, ref):
    dn = DefinedName(nom, attr_text=ref)
    try:
        wb.defined_names[nom] = dn          # openpyxl >= 3.1
    except TypeError:
        wb.defined_names.append(dn)         # openpyxl 3.0


# ======================================================================
# 3. CONSTRUCTION DU CLASSEUR
# ======================================================================

S_LIS, S_REC, S_PAR = "Lisez-moi", "Récap", "Paramètres"
S_CO, S_B, S_SP = "0 Consommation", "1a Lien prix-volume", "1b Spot 2027"
S_AL, S_ST, S_M1 = "1c Prix T2 anticipé", "1d Stratégies", "1e Marges T1"
S_T2, S_T3, S_R4 = "2 Décision T2", "3 Position 01-01", "4 Résultat 01-01"
S_H, S_BL = "Heures 2027", "Blocs"


def strategies_excel(D):
    """Stratégies du tableau de bord -> ratios T1 (8) et cibles T2 (8)."""
    PQ = [0, 0, 0, 1, 1, 1, 2, 3]
    out = []
    for s in D["strategies"]:
        r1, r2, actif = [0.0] * 8, [0.0] * 8, 0
        for e in s.get("etapes", []):
            t, r = e.get("type"), e.get("ratio")
            if e.get("date") == "T1" and t and t != "aucun":
                if t == "mois":
                    r1 = [float(x) for x in r]
                elif t == "trimestres":
                    r1 = [float(r[PQ[p]]) for p in range(8)]
                else:
                    r1 = [float(r if not isinstance(r, list) else r[0])] * 8
                if t != "mois":
                    print(f"  Attention : {s['nom']}, étape T1 « {t} » convertie en mois (approximation).")
            elif e.get("date") == "T2" and t == "complement":
                r2, actif = [float(x) for x in r], 1
                if e.get("prod", "mois") != "mois":
                    print(f"  Attention : {s['nom']}, complément en {e.get('prod')} traité en mois.")
            elif e.get("date") == "T2" and t and t != "aucun":
                print(f"  Attention : {s['nom']}, étape T2 « {t} » non reprise dans l'Excel.")
        out.append({"nom": s["nom"], "r1": r1, "r2": r2, "actif": actif})
    return out


def construire_classeur(D, H, J, feries, annees):
    d = D["defaults"]
    dates = d["dates"]
    strats = strategies_excel(D)
    NS = len(strats)
    NY = len(annees)
    if NY != 4:
        raise SystemExit("Le classeur est prévu pour 4 années de scénarios (config.ANNEES_SCENARIOS).")
    N = len(H)

    wb = Workbook()
    try:
        wb._named_styles["Normal"].font = Font(name=POLICE, size=10)
    except Exception:
        pass
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    ws_lis = wb.active
    ws_lis.title = S_LIS
    ws_rec = wb.create_sheet(S_REC)
    ws_par = wb.create_sheet(S_PAR)
    ws_co = wb.create_sheet(S_CO)
    ws_b = wb.create_sheet(S_B)
    ws_sp = wb.create_sheet(S_SP)
    ws_al = wb.create_sheet(S_AL)
    ws_st = wb.create_sheet(S_ST)
    ws_m1 = wb.create_sheet(S_M1)
    ws_t2 = wb.create_sheet(S_T2)
    ws_t3 = wb.create_sheet(S_T3)
    ws_r4 = wb.create_sheet(S_R4)
    ws_h = wb.create_sheet(S_H)
    ws_bl = wb.create_sheet(S_BL)

    # ------------------------------------------------------------------
    # PARAMÈTRES
    # ------------------------------------------------------------------
    ws = ws_par
    titre(ws, "Paramètres du modèle",
          "Cellules bleues sur fond jaune : hypothèses que l'on peut modifier. Tout le classeur se recalcule.")
    section(ws, 4, "Portefeuille de clients", 1, 5)
    entetes(ws, 5, 1, ["Paramètre", "T1 (22/09/2026)", "T2 (01/12/2026)", "T3 (31/12/2026)", "Source"])
    lignes = [("Clients RES1", "nbRes"), ("Clients PRO1", "nbPro")]
    for i, (lib, k) in enumerate(lignes):
        r = 6 + i
        ecrire(ws, r, 1, lib)
        for j, t in enumerate(["T1", "T2", "T3"]):
            if t in dates:
                entree(ws, r, 2 + j, dates[t][k], "#,##0")
        ecrire(ws, r, 5, "config.py (fichiers Moodle pour T2 et T3)", font=F_SOUS)
    ecrire(ws, 8, 1, "θ RES1 (facteur d'utilisation)"); entree(ws, 8, 2, d["thRes"], "0.00000")
    ecrire(ws, 9, 1, "θ PRO1"); entree(ws, 9, 2, d["thPro"], "0.00000")
    ecrire(ws, 10, 1, "Puissance souscrite RES1 (kVA)"); entree(ws, 10, 2, d["pRes"])
    ecrire(ws, 11, 1, "Puissance souscrite PRO1 (kVA)"); entree(ws, 11, 2, d["pPro"])
    ecrire(ws, 12, 1, "Puissance moyenne RES1 = N × θ × kVA (MW)")
    ecrire(ws, 13, 1, "Puissance moyenne PRO1 = N × θ × kVA (MW)")
    ecrire(ws, 14, 1, "Rapport au portefeuille du 22/09, RES1")
    ecrire(ws, 15, 1, "Rapport au portefeuille du 22/09, PRO1")
    for j in range(3):
        c = 2 + j
        ecrire(ws, 12, c, f"={L(c)}6*$B$8*$B$10/1000", "0.000")
        ecrire(ws, 13, c, f"={L(c)}7*$B$9*$B$11/1000", "0.000")
        ecrire(ws, 14, c, f"={L(c)}12/$B$12", "0.0000")
        ecrire(ws, 15, c, f"={L(c)}13/$B$13", "0.0000")
    for nom, ref in [("pmoy_res_t1", "$B$12"), ("pmoy_pro_t1", "$B$13"), ("pmoy_res_t3", "$D$12"),
                     ("pmoy_pro_t3", "$D$13"), ("k_res_t2", "$C$14"), ("k_pro_t2", "$C$15")]:
        nommer(wb, nom, f"{q(S_PAR)}!{ref}")

    section(ws, 17, "Règles du Board et hypothèses de marché", 1, 5)
    params = [
        (18, "Prix de vente (€/MWh)", d["prixVente"], "0.00", "prix_vente", "guidelines"),
        (19, "Marge cible (€/MWh)", d.get("margeCible", 10.0), "0.00", "marge_cible", "guidelines"),
        (20, "Plancher du pire cas (€/MWh)", d["plancher"], "0.00", "plancher", "Board"),
        (21, "Prime de risque (spot attendu = forward × (1 − prime))", d["prime"], "0.0%", "prime", "guidelines"),
        (22, "Alpha (part du scénario intégrée par le marché au 01/12)", d["alpha"], "0.00", "alpha", "hypothèse justifiée en 1c"),
        (26, "Recentrage des anomalies (1 = oui, 0 = non)", 1 if d.get("recentrer", True) else 0, "0", "recentrage", "voir 0 Consommation"),
        (27, "Seuil de chauffage (°C)", D["meta"].get("seuil", 15.0), "0.0", "seuil", "cours"),
    ]
    for r, lib, v, f, nom, src in params:
        ecrire(ws, r, 1, lib)
        entree(ws, r, 2, v, f)
        ecrire(ws, r, 5, src, font=F_SOUS)
        nommer(wb, nom, f"{q(S_PAR)}!$B${r}")
    ecrire(ws, 23, 1, "b calculé par la régression (feuille 1a)")
    ecrire(ws, 23, 2, "=" + A(S_B, 2, 5), "0.000")
    ecrire(ws, 24, 1, "b imposé (laisser vide pour utiliser b calculé)")
    entree(ws, 24, 2, None, "0.000")
    ecrire(ws, 25, 1, "b utilisé dans le modèle")
    resultat(ws, 25, 2, "=IF(ISBLANK(B24),ROUND(B23,3),B24)", "0.000")
    nommer(wb, "b_util", f"{q(S_PAR)}!$B$25")

    section(ws, 29, "Tests de stress (crise de prix)", 1, 5)
    entetes(ws, 30, 1, ["Choc", "Hausse du spot", "Compte dans le pire cas (1/0)"])
    chocs = list(d.get("chocs", []))[:2]
    while len(chocs) < 2:
        chocs.append(["(aucun)", 0.0])
    compte = 1 if d.get("stress", True) else 0
    for i, (nom, v) in enumerate(chocs):
        r = 31 + i
        ecrire(ws, r, 1, nom)
        entree(ws, r, 2, float(v), "0%")
        entree(ws, r, 3, compte if float(v) != 0 else 0, "0")
    nommer(wb, "choc_tension", f"{q(S_PAR)}!$B$31")
    nommer(wb, "choc_crise", f"{q(S_PAR)}!$B$32")
    nommer(wb, "compte_tension", f"{q(S_PAR)}!$C$31")
    nommer(wb, "compte_crise", f"{q(S_PAR)}!$C$32")

    section(ws, 34, "Probabilités des scénarios météo", 1, 5)
    entetes(ws, 35, 1, ["Année météo", "Probabilité"])
    probas = d.get("proba", [1 / NY] * NY)
    for i, a in enumerate(annees):
        ecrire(ws, 36 + i, 1, a)
        entree(ws, 36 + i, 2, float(probas[i]), "0%")
    PROBAS = RG(S_PAR, 2, 36, 2, 39)

    section(ws, 41, "Volatilité des prix (probabilité des crises)", 1, 5)
    vol_params = [
        (42, "Volatilité annuelle des forwards (CRE)", d.get("mc", {}).get("sigma", 0.20), "0%", "volatilite"),
        (43, "Horizon vu du 22/09 (années jusqu'au milieu de 2027)", 0.78, "0.00", "horizon_t1"),
        (44, "Horizon vu du 01/12 (années)", 0.58, "0.00", "horizon_t2"),
        (45, "Jours du 22/09 au 01/12", 70, "0", "jours_t1_t2"),
    ]
    for r, lib, v, f, nom in vol_params:
        ecrire(ws, r, 1, lib)
        entree(ws, r, 2, v, f)
        nommer(wb, nom, f"{q(S_PAR)}!$B${r}")

    section(ws, 47, "Prix EEX (€/MWh)", 1, 7)
    entetes(ws, 48, 1, ["Produit", "Base 22/09", "Peak 22/09", "Base 01/12", "Peak 01/12", "Base 31/12", "Peak 31/12"])
    produits = ["Cal-27", "Q1-27", "Q2-27", "Q3-27", "Q4-27"] + [f"{m}-27" for m in CLES_MOIS]
    for i, prod in enumerate(produits):
        r = 49 + i
        ecrire(ws, r, 1, prod)
        for j, t in enumerate(["T1", "T2", "T3"]):
            if t not in dates:
                continue
            dd = dates[t]
            if i == 0:
                vb, vp = dd["calBase"], dd["calPeak"]
            elif i <= 4:
                vb, vp = dd["q"][i - 1]
            else:
                m = {x[0]: x for x in dd.get("mois", [])}.get(CLES_MOIS[i - 5])
                vb, vp = (m[1], m[2]) if m else (None, None)
            entree(ws, r, 2 + 2 * j, vb, "0.00")
            entree(ws, r, 3 + 2 * j, vp, "0.00")
    ROW_PRIX = {p: 49 + i for i, p in enumerate(produits)}   # Cal 49, Q1 50..Q4 53, Jan 54..Jun 59
    largeurs(ws, [58, 16, 16, 16, 16, 16, 16])

    def prix_periode(p, col):
        """Formule du prix d'une période (mois s'il est coté, sinon trimestre)."""
        if p <= 6:
            rm, rq = ROW_PRIX[f"{CLES_MOIS[p - 1]}-27"], ROW_PRIX[f"Q{(p - 1) // 3 + 1}-27"]
            cm, cq = A(S_PAR, col, rm), A(S_PAR, col, rq)
            return f"=IF(ISBLANK({cm}),{cq},{cm})"
        return "=" + A(S_PAR, col, ROW_PRIX["Q3-27" if p == 7 else "Q4-27"])

    # ------------------------------------------------------------------
    # HEURES 2027 (8 760 lignes)
    # ------------------------------------------------------------------
    ws = ws_h
    titre(ws, "Données horaires 2027 : consommation, météo et prix spot",
          "Fond bleu clair : données préparées par Python. Colonnes blanches : formules Excel.")
    note(ws, 3, 1, "Une ligne par heure de 2027. Le coefficient de profil est déjà dé-météorisé "
                   "et aligné sur le calendrier 2027 (jours de semaine et jours fériés).")
    H0 = 6
    HN = H0 + N - 1
    C_DT, C_MOIS, C_PER, C_PK, C_BLOC = 1, 2, 3, 4, 5
    C_CR, C_CP, C_LR, C_LP, C_LT, C_TN, C_GR, C_GP = 6, 7, 8, 9, 10, 11, 12, 13
    SC = lambda i: 14 + 9 * i          # première colonne du scénario i
    ent = ["Date et heure", "Mois", "Période (1-6 = mois, 7 = Q3, 8 = Q4)", "Peak (1 = 8h-20h lun-ven)",
           "Bloc (période × peak)", "Coefficient profil RES1 (normal)", "Coefficient profil PRO1 (normal)",
           "Conso normale RES1 (MW)", "Conso normale PRO1 (MW)", "Conso normale totale (MW)",
           "Température normale (°C)", "Gradient RES1 (/°C)", "Gradient PRO1 (/°C)"]
    for i, a in enumerate(annees):
        ent += [f"Écart de température {a} (°C)", f"T réelle {a} (°C)", f"CM RES1 {a}", f"CM PRO1 {a}",
                f"Conso RES1 météo {a} (MW)", f"Conso PRO1 météo {a} (MW)", f"Prix historique {a} (€/MWh)",
                f"Forme {a} (moyenne 1 par bloc)", f"Spot 2027 scénario {a} (€/MWh)"]
    entetes(ws, 5, 1, ent, 60)
    dts = H.index.tz_localize(None).to_pydatetime()
    cols_data = {C_CR: "coef_res1", C_CP: "coef_pro1", C_TN: "tn", C_GR: "g_res1", C_GP: "g_pro1"}
    fm_data = {C_CR: "coef", C_CP: "coef", C_TN: "temp", C_GR: "0.00000", C_GP: "0.00000"}
    vals = {k: H[v].to_numpy() for k, v in cols_data.items()}
    dtv = {i: H[f"dt_{a}"].to_numpy() for i, a in enumerate(annees)}
    prv = {i: H[f"prix_{a}"].to_numpy() for i, a in enumerate(annees)}

    # plages des niveaux de prix (feuille 1b) et des prix moyens par bloc (feuille Blocs)
    SP_NIV = lambda i: RG(S_SP, 9 + i, 16, 9 + i, 31)
    BL_PMOY = lambda i: RG(S_BL, 8 + 6 * i + 5, 5, 8 + 6 * i + 5, 20)

    for k in range(N):
        r = H0 + k
        donnee(ws, r, C_DT, dts[k], "dt")
        ws.cell(row=r, column=C_MOIS, value=f"=MONTH(A{r})")
        ws.cell(row=r, column=C_PER, value=f"=IF(B{r}<=6,B{r},IF(B{r}<=9,7,8))")
        ws.cell(row=r, column=C_PK, value=f"=IF(AND(WEEKDAY(A{r},2)<=5,HOUR(A{r})>=8,HOUR(A{r})<20),1,0)")
        ws.cell(row=r, column=C_BLOC, value=f"=2*(C{r}-1)+D{r}+1")
        for c, v in vals.items():
            donnee(ws, r, c, float(v[k]), fm_data[c])
        ws.cell(row=r, column=C_LR, value=f"=F{r}*pmoy_res_t1").number_format = FMT["mw"]
        ws.cell(row=r, column=C_LP, value=f"=G{r}*pmoy_pro_t1").number_format = FMT["mw"]
        ws.cell(row=r, column=C_LT, value=f"=H{r}+I{r}").number_format = FMT["mw"]
        for i in range(NY):
            c0 = SC(i)
            dt_, tr, cmr, cmp_, lr, lp, px, fo, sp = [L(c0 + j) for j in range(9)]
            donnee(ws, r, c0, float(dtv[i][k]), "temp")
            ws.cell(row=r, column=c0 + 1, value=f"=K{r}+{dt_}{r}").number_format = FMT["temp"]
            for cc, g in ((c0 + 2, "L"), (c0 + 3, "M")):
                ws.cell(row=r, column=cc, value=(
                    f"=IF(AND({tr}{r}<seuil,K{r}<seuil),1+{g}{r}*(K{r}-{tr}{r}),"
                    f"IF(AND({tr}{r}<seuil,K{r}>=seuil),1+{g}{r}*(seuil-{tr}{r}),"
                    f"IF(AND(K{r}<seuil,{tr}{r}>=seuil),1+{g}{r}*(K{r}-seuil),1)))")).number_format = FMT["coef"]
            ws.cell(row=r, column=c0 + 4, value=f"=H{r}*{cmr}{r}").number_format = FMT["mw"]
            ws.cell(row=r, column=c0 + 5, value=f"=I{r}*{cmp_}{r}").number_format = FMT["mw"]
            donnee(ws, r, c0 + 6, float(prv[i][k]), "prix")
            ws.cell(row=r, column=c0 + 7, value=f"={px}{r}/INDEX({BL_PMOY(i)},E{r})").number_format = FMT["coef"]
            ws.cell(row=r, column=c0 + 8, value=f"=INDEX({SP_NIV(i)},E{r})*{fo}{r}").number_format = FMT["prix"]
    ws.freeze_panes = "B6"
    largeurs(ws, [17] + [9] * 4 + [12] * (SC(NY) - 6))

    def HC(c):
        """Plage complète d'une colonne de la feuille horaire."""
        return RG(S_H, c, H0, c, HN)

    # ------------------------------------------------------------------
    # BLOCS (agrégats par période × peak)
    # ------------------------------------------------------------------
    ws = ws_bl
    titre(ws, "Blocs : sommes horaires par période et par peak / hors peak",
          "Feuille technique. Les marges se calculent bloc par bloc : le niveau de prix est constant dans un bloc.")
    ent = ["Bloc", "Période", "Nom", "Peak", "Heures", "Σ conso normale RES1 (MWh)", "Σ conso normale PRO1 (MWh)"]
    for a in annees:
        ent += [f"Σ RES1 météo {a}", f"Σ PRO1 météo {a}", f"Σ RES1 × forme {a}", f"Σ PRO1 × forme {a}",
                f"Σ prix historique {a}", f"Prix moyen du bloc {a}"]
    entetes(ws, 4, 1, ent, 45)
    for b in range(1, 17):
        r = 4 + b
        p, k = (b - 1) // 2 + 1, (b - 1) % 2
        ecrire(ws, r, 1, b)
        ecrire(ws, r, 2, p)
        ecrire(ws, r, 3, NOMS_P[p - 1])
        ecrire(ws, r, 4, k)
        ecrire(ws, r, 5, f"=COUNTIF({HC(C_BLOC)},A{r})", "int")
        ecrire(ws, r, 6, f"=SUMIFS({HC(C_LR)},{HC(C_BLOC)},A{r})", "mwh")
        ecrire(ws, r, 7, f"=SUMIFS({HC(C_LP)},{HC(C_BLOC)},A{r})", "mwh")
        for i in range(NY):
            c0, h0 = 8 + 6 * i, SC(i)
            ecrire(ws, r, c0, f"=SUMIFS({HC(h0 + 4)},{HC(C_BLOC)},A{r})", "mwh")
            ecrire(ws, r, c0 + 1, f"=SUMIFS({HC(h0 + 5)},{HC(C_BLOC)},A{r})", "mwh")
            ecrire(ws, r, c0 + 2, f"=SUMPRODUCT(({HC(C_BLOC)}=A{r})*{HC(h0 + 4)}*{HC(h0 + 7)})", "mwh")
            ecrire(ws, r, c0 + 3, f"=SUMPRODUCT(({HC(C_BLOC)}=A{r})*{HC(h0 + 5)}*{HC(h0 + 7)})", "mwh")
            ecrire(ws, r, c0 + 4, f"=SUMIFS({HC(h0 + 6)},{HC(C_BLOC)},A{r})", "eur")
            ecrire(ws, r, c0 + 5, f"={L(c0 + 4)}{r}/E{r}", "prix")
    section(ws, 22, "Heures par période", 1, 7)
    entetes(ws, 23, 1, ["Période", "Nom", "Heures", "Heures peak", "Heures hors peak", "Trimestre"])
    for p in range(1, 9):
        r = 23 + p
        ro, rp = 5 + 2 * (p - 1), 6 + 2 * (p - 1)
        ecrire(ws, r, 1, p)
        ecrire(ws, r, 2, NOMS_P[p - 1])
        ecrire(ws, r, 3, f"=E{ro}+E{rp}", "int")
        ecrire(ws, r, 4, f"=E{rp}", "int")
        ecrire(ws, r, 5, f"=E{ro}", "int")
        ecrire(ws, r, 6, (p - 1) // 3 + 1 if p <= 6 else p - 4)
    largeurs(ws, [7, 8, 16, 6, 8] + [13] * 30)
    HP = RG(S_BL, 3, 24, 3, 31)       # heures par période
    HPK = RG(S_BL, 4, 24, 4, 31)      # heures peak par période
    BL = lambda c: RG(S_BL, c, 5, c, 20)

    def bl_ro(p):
        return 5 + 2 * (p - 1)

    # ------------------------------------------------------------------
    # 0 CONSOMMATION
    # ------------------------------------------------------------------
    ws = ws_co
    titre(ws, "Étape 0 : consommation 2027 (portefeuille du 22/09)",
          "Conso = coefficient de profil × N × θ × kVA, puis × CM pour chaque année météo (feuille Heures 2027).")
    section(ws, 3, "Volumes mensuels (MWh) et partage peak / hors peak", 1, 9)
    entetes(ws, 4, 1, ["Mois", "Volume normal", "dont peak", "dont hors peak", "Part peak"]
            + [f"Météo {a}" for a in annees], 32)
    mois_noms = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août",
                 "Septembre", "Octobre", "Novembre", "Décembre"]
    for m in range(1, 13):
        r = 4 + m
        ecrire(ws, r, 1, mois_noms[m - 1])
        ecrire(ws, r, 2, f"=SUMIFS({HC(C_LT)},{HC(C_MOIS)},{m})", "mwh")
        ecrire(ws, r, 3, f"=SUMIFS({HC(C_LT)},{HC(C_MOIS)},{m},{HC(C_PK)},1)", "mwh")
        ecrire(ws, r, 4, f"=B{r}-C{r}", "mwh")
        ecrire(ws, r, 5, f"=C{r}/B{r}", "pct")
        for i in range(NY):
            h0 = SC(i)
            ecrire(ws, r, 6 + i, f"=SUMIFS({HC(h0 + 4)},{HC(C_MOIS)},{m})+SUMIFS({HC(h0 + 5)},{HC(C_MOIS)},{m})", "mwh")
    ecrire(ws, 17, 1, "Année", bold=True)
    for c in range(2, 6 + NY):
        if c == 5:
            resultat(ws, 17, c, "=C17/B17", "pct")
        else:
            resultat(ws, 17, c, f"=SUM({L(c)}5:{L(c)}16)", "mwh")
    note(ws, 18, 1, "Le volume normal de l'année est le « 100 % » de tous les taux de couverture.")

    section(ws, 19, "Volume annuel de chaque scénario météo", 1, 9)
    entetes(ws, 20, 1, ["Scénario", "Volume (MWh)", "Écart à la normale", "Anomalie recentrée"])
    for i, a in enumerate(annees):
        r = 21 + i
        ecrire(ws, r, 1, f"Météo {a}")
        ecrire(ws, r, 2, f"={L(6 + i)}17", "mwh")
        ecrire(ws, r, 3, f"=B{r}/$B$17-1", "0.00%")
        ecrire(ws, r, 4, f"=C{r}-AVERAGE($C$21:$C$24)*recentrage", "0.00%")
    note(ws, 25, 1, "Recentrage : les 4 années ont été plus douces que la normale ; on mesure chaque année "
                    "par rapport à leur moyenne pour que le prix moyen reste ancré sur le forward.")

    section(ws, 26, "Couverture 100 % par période : Base = conso moyenne hors peak, Peak = supplément en peak", 1, 10)
    entetes(ws, 27, 1, ["Période", "Produit", "Heures", "Heures peak", "Énergie normale (MWh)",
                        "Conso moyenne hors peak (MW)", "Conso moyenne peak (MW)", "Base 100 % (MW)",
                        "Peak 100 % (MW)", "Contrôle : énergie couverte (MWh)"], 45)
    for p in range(1, 9):
        r = 27 + p
        ro, rp = bl_ro(p), bl_ro(p) + 1
        ecrire(ws, r, 1, NOMS_P[p - 1])
        ecrire(ws, r, 2, PROD_P[p - 1])
        ecrire(ws, r, 3, "=" + A(S_BL, 3, 23 + p), "int")
        ecrire(ws, r, 4, "=" + A(S_BL, 4, 23 + p), "int")
        ecrire(ws, r, 5, f"={A(S_BL, 6, ro)}+{A(S_BL, 7, ro)}+{A(S_BL, 6, rp)}+{A(S_BL, 7, rp)}", "mwh")
        ecrire(ws, r, 6, f"=({A(S_BL, 6, ro)}+{A(S_BL, 7, ro)})/{A(S_BL, 5, ro)}", "mw")
        ecrire(ws, r, 7, f"=({A(S_BL, 6, rp)}+{A(S_BL, 7, rp)})/{A(S_BL, 5, rp)}", "mw")
        ecrire(ws, r, 8, f"=IF(G{r}>=F{r},F{r},E{r}/C{r})", "mw")
        ecrire(ws, r, 9, f"=IF(G{r}>=F{r},G{r}-F{r},0)", "mw")
        ecrire(ws, r, 10, f"=H{r}*C{r}+I{r}*D{r}", "mwh")
    resultat(ws, 36, 5, "=SUM(E28:E35)", "mwh")
    ecrire(ws, 36, 1, "Total", bold=True)

    section(ws, 37, "Anomalies de volume par période (prix spot et prix T2 anticipé)", 1, 14)
    entetes(ws, 38, 1, ["Période"] + [f"Brute {a}" for a in annees] + [f"Recentrée {a}" for a in annees]
            + [f"Énergie météo {a} (MWh)" for a in annees], 32)
    for p in range(1, 9):
        r = 38 + p
        ro, rp = bl_ro(p), bl_ro(p) + 1
        ecrire(ws, r, 1, NOMS_P[p - 1])
        for i in range(NY):
            c0 = 8 + 6 * i
            ecrire(ws, r, 10 + i, f"={A(S_BL, c0, ro)}+{A(S_BL, c0 + 1, ro)}+{A(S_BL, c0, rp)}+{A(S_BL, c0 + 1, rp)}", "mwh")
            ecrire(ws, r, 2 + i, f"={L(10 + i)}{r}/$E${27 + p}-1", "0.00%")
            ecrire(ws, r, 6 + i, f"={L(2 + i)}{r}-AVERAGE($B{r}:$E{r})*recentrage", "0.00%")
    section(ws, 48, "Anomalies de volume par trimestre", 1, 14)
    entetes(ws, 49, 1, ["Trimestre", "Énergie normale (MWh)"] + [f"Brute {a}" for a in annees]
            + [f"Recentrée {a}" for a in annees] + [f"Énergie météo {a}" for a in annees], 32)
    for qq in range(1, 5):
        r = 49 + qq
        ps = PERIODES_Q[qq]
        ecrire(ws, r, 1, f"Q{qq}")
        ecrire(ws, r, 2, f"=SUM(E{27 + ps[0]}:E{27 + ps[-1]})", "mwh")
        for i in range(NY):
            cc = L(10 + i)
            ecrire(ws, r, 11 + i, f"=SUM({cc}{38 + ps[0]}:{cc}{38 + ps[-1]})", "mwh")
            ecrire(ws, r, 3 + i, f"={L(11 + i)}{r}/$B{r}-1", "0.00%")
            ecrire(ws, r, 7 + i, f"={L(3 + i)}{r}-AVERAGE($C{r}:$F{r})*recentrage", "0.00%")
    largeurs(ws, [20, 14, 13, 13, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14])
    ch = BarChart()
    ch.type = "col"
    ch.title = "Volumes mensuels 2027 (MWh)"
    ch.y_axis.title = "MWh"
    ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=16), titles_from_data=True)
    for i in range(NY):
        ch.add_data(Reference(ws, min_col=6 + i, min_row=4, max_row=16), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=16))
    ch.height, ch.width = 9, 20
    ws.add_chart(ch, "L3")
    ANOM_P = lambda i: RG(S_CO, 6 + i, 39, 6 + i, 46)     # anomalie recentrée par période, scénario i

    # ------------------------------------------------------------------
    # 1a LIEN PRIX-VOLUME
    # ------------------------------------------------------------------
    ws = ws_b
    titre(ws, "Étape 1a : lien prix-volume (pente b)",
          "x = anomalie de volume du jour, y = anomalie de prix du jour, sur les jours d'hiver (novembre à mars).")
    nJ = len(J)
    A0 = 22
    AN = A0 + nJ - 1
    hiver = [i for i in range(nJ) if J["date"][i].month in getattr(config, "MOIS_HIVER", (11, 12, 1, 2, 3))]
    B0 = 22
    BN = B0 + len(hiver) - 1
    # colonnes du tableau d'hiver
    CB_D, CB_A, CB_X, CB_Y, CB_XC = 18, 19, 20, 21, 22
    XB, YB, XCB = (RG(S_B, c, B0, c, BN) for c in (CB_X, CB_Y, CB_XC))
    rows_annee = {}
    for jj, i in enumerate(hiver):
        rows_annee.setdefault(J["date"][i].year, []).append(B0 + jj)

    section(ws, 3, "Résultats de la régression y = a + b x (moindres carrés)", 1, 6)
    res_lignes = [
        (5, "b = Cov(x, y) / Var(x)", f"=COVAR({XB},{YB})/VARP({XB})", "b"),
        (6, "Contrôle : pente Excel (SLOPE)", f"=SLOPE({YB},{XB})", "b"),
        (7, "a (ordonnée à l'origine)", f"=INTERCEPT({YB},{XB})", "0.0000"),
        (8, "R²", f"=RSQ({YB},{XB})", "0.00"),
        (9, "Nombre de jours d'hiver", f"=COUNT({XB})", "int"),
    ]
    for r, lib, f, fm in res_lignes:
        ecrire(ws, r, 1, lib)
        (resultat if r in (5, 8) else ecrire)(ws, r, 2, f, fm)
    ecrire(ws, 10, 1, "b calculé par Python (contrôle)")
    ecrire(ws, 10, 2, float(D["meta"]["b"]), "b")
    ecrire(ws, 11, 1, "b par année :", bold=True)
    for k, (an, rr) in enumerate(sorted(rows_annee.items())):
        x_ = RG(S_B, CB_X, rr[0], CB_X, rr[-1])
        y_ = RG(S_B, CB_Y, rr[0], CB_Y, rr[-1])
        ecrire(ws, 12 + k, 1, f"   hiver {an}")
        ecrire(ws, 12 + k, 2, f"=COVAR({x_},{y_})/VARP({x_})", "0.00")
    r_sans = 12 + len(rows_annee)
    autres = [r for an, rr in rows_annee.items() if an != 2022 for r in rr]
    if autres and 2022 in rows_annee:
        x_ = RG(S_B, CB_X, min(autres), CB_X, max(autres))
        y_ = RG(S_B, CB_Y, min(autres), CB_Y, max(autres))
        ecrire(ws, r_sans, 1, "b sans 2022 (année de crise)")
        ecrire(ws, r_sans, 2, f"=COVAR({x_},{y_})/VARP({x_})", "0.00")
    ecrire(ws, 5, 4, "b corrigé : x mesuré, comme y, par rapport à la moyenne de son mois et type de jour", bold=True)
    ecrire(ws, 6, 4, "b corrigé")
    resultat(ws, 6, 5, f"=COVAR({XCB},{YB})/VARP({XCB})", "b")
    ecrire(ws, 7, 4, "R² corrigé")
    ecrire(ws, 7, 5, f"=RSQ({YB},{XCB})", "0.00")
    note(ws, 8, 4, "Le prix est comparé à la moyenne de son mois, le volume à la normale : un mois entier plus doux")
    note(ws, 9, 4, "décale ses points sans changer les prix, ce qui aplatit un peu la droite. b corrigé sert de test de sensibilité.")
    note(ws, 10, 4, "Pour l'utiliser dans tout le modèle : le taper dans Paramètres, « b imposé ».")

    section(ws, 20, "Tous les jours 2022-2025", 1, 13)
    entetes(ws, 21, 1, ["Date", "Année", "Mois", "Férié", "Type de jour", "Conso normale (MW)",
                        "Conso avec la météo réelle (MW)", "x = anomalie de volume", "Prix du jour (€/MWh)",
                        "Prix de référence (même mois, même type de jour)", "y = anomalie de prix",
                        "Hiver (1 = nov. à mars)", "x recentré (mois, type de jour)"], 58)
    ecrire(ws, 21, 15, "Jours fériés", bold=True)
    for k, fd in enumerate(feries):
        ecrire(ws, 22 + k, 15, fd, "date")
    FER = f"$O$22:$O${21 + len(feries)}"
    mois_hiver = getattr(config, "MOIS_HIVER", (11, 12, 1, 2, 3))
    cond_hiver = "OR(" + ",".join(f"C{{r}}={m}" for m in mois_hiver) + ")"
    for k in range(nJ):
        r = A0 + k
        donnee(ws, r, 1, J["date"][k], "date")
        ecrire(ws, r, 2, f"=YEAR(A{r})")
        ecrire(ws, r, 3, f"=MONTH(A{r})")
        ecrire(ws, r, 4, f"=IF(COUNTIF({FER},A{r})>0,1,0)")
        ecrire(ws, r, 5, f'=IF(OR(D{r}=1,WEEKDAY(A{r},2)=7),"dim_ferie",IF(WEEKDAY(A{r},2)=6,"samedi","ouvre"))')
        donnee(ws, r, 6, float(J["load_normal_total_mw"][k]), "mw")
        donnee(ws, r, 7, float(J["load_final_total_mw"][k]), "mw")
        ecrire(ws, r, 8, f"=G{r}/F{r}-1", "0.00%")
        donnee(ws, r, 9, float(J["prix_jour"][k]), "prix")
        ecrire(ws, r, 10, f"=AVERAGEIFS($I${A0}:$I${AN},$B${A0}:$B${AN},B{r},$C${A0}:$C${AN},C{r},$E${A0}:$E${AN},E{r})", "prix")
        ecrire(ws, r, 11, f"=I{r}/J{r}-1", "0.00%")
        ecrire(ws, r, 12, "=IF(" + cond_hiver.format(r=r) + ",1,0)")
        ecrire(ws, r, 13, f"=H{r}-AVERAGEIFS($H${A0}:$H${AN},$B${A0}:$B${AN},B{r},$C${A0}:$C${AN},C{r},$E${A0}:$E${AN},E{r})", "0.00%")
    section(ws, 20, "Jours d'hiver retenus pour la régression", CB_D, CB_XC)
    entetes(ws, 21, CB_D, ["Date", "Année", "x", "y", "x recentré"], 58)
    for jj, i in enumerate(hiver):
        r, ra = B0 + jj, A0 + i
        ecrire(ws, r, CB_D, f"=A{ra}", "date")
        ecrire(ws, r, CB_A, f"=B{ra}")
        ecrire(ws, r, CB_X, f"=H{ra}", "0.00%")
        ecrire(ws, r, CB_Y, f"=K{ra}", "0.00%")
        ecrire(ws, r, CB_XC, f"=M{ra}", "0.00%")
    largeurs(ws, [46, 12, 7, 7, 30, 12, 13, 12, 11, 14, 12, 9, 12, 3, 12, 3, 3, 12, 8, 10, 10, 10])
    sc = ScatterChart()
    sc.title = "Anomalie de prix (y) en fonction de l'anomalie de volume (x), jours d'hiver"
    sc.style = 13
    sc.x_axis.title = "Anomalie de volume"
    sc.y_axis.title = "Anomalie de prix"
    ser = Series(Reference(ws, min_col=CB_Y, min_row=B0, max_row=BN),
                 Reference(ws, min_col=CB_X, min_row=B0, max_row=BN), title="Jours d'hiver")
    ser.marker.symbol = "circle"
    ser.marker.size = 4
    ser.graphicalProperties.line.noFill = True
    ser.trendline = Trendline(trendlineType="linear", dispRSqr=True, dispEq=True)
    sc.series.append(ser)
    sc.height, sc.width = 10, 18
    ws.add_chart(sc, "G3")

    # ------------------------------------------------------------------
    # 1b SPOT 2027
    # ------------------------------------------------------------------
    ws = ws_sp
    titre(ws, "Étape 1b : scénarios de prix spot 2027 (prix du 22/09)",
          "Spot de l'heure = niveau du bloc × forme historique ; niveau = forward × (1 − prime) × (1 + b × anomalie).")
    section(ws, 3, "Prix EEX par période (le mois s'il est coté, sinon le trimestre)", 1, 7)
    entetes(ws, 4, 1, ["Période", "Produit", "Forward Base", "Forward Peak", "Heures", "Heures peak",
                       "Hors peak déduit = (Base × H − Peak × Hpk) / (H − Hpk)"], 45)
    for p in range(1, 9):
        r = 4 + p
        ecrire(ws, r, 1, NOMS_P[p - 1])
        ecrire(ws, r, 2, PROD_P[p - 1])
        ecrire(ws, r, 3, prix_periode(p, 2), "prix")
        ecrire(ws, r, 4, prix_periode(p, 3), "prix")
        ecrire(ws, r, 5, "=" + A(S_BL, 3, 23 + p), "int")
        ecrire(ws, r, 6, "=" + A(S_BL, 4, 23 + p), "int")
        ecrire(ws, r, 7, f"=(C{r}*E{r}-D{r}*F{r})/(E{r}-F{r})", "prix")
    FB1, FP1, FO1 = RG(S_SP, 3, 5, 3, 12), RG(S_SP, 4, 5, 4, 12), RG(S_SP, 7, 5, 7, 12)

    section(ws, 14, "Niveau de prix de chaque bloc et de chaque scénario (€/MWh)", 1, 12)
    entetes(ws, 15, 1, ["Bloc", "Période", "Peak", "Forward du bloc"] + [f"Anomalie {a}" for a in annees]
            + [f"Niveau {a}" for a in annees], 32)
    for b in range(1, 17):
        r = 15 + b
        p, k = (b - 1) // 2 + 1, (b - 1) % 2
        ecrire(ws, r, 1, b)
        ecrire(ws, r, 2, p)
        ecrire(ws, r, 3, k)
        ecrire(ws, r, 4, f"=IF(C{r}=1,INDEX($D$5:$D$12,B{r}),INDEX($G$5:$G$12,B{r}))", "prix")
        for i in range(NY):
            ecrire(ws, r, 5 + i, f"=INDEX({ANOM_P(i)},B{r})", "0.00%")
            ecrire(ws, r, 9 + i, f"=$D{r}*(1-prime)*(1+b_util*{L(5 + i)}{r})", "prix")

    section(ws, 33, "Spot 2027 moyen (€/MWh)", 1, 7)
    entetes(ws, 34, 1, ["Scénario", "Année", "Q1", "Q2", "Q3", "Q4"])
    for i, a in enumerate(annees):
        r = 35 + i
        col = HC(SC(i) + 8)
        ecrire(ws, r, 1, f"Météo {a}")
        resultat(ws, r, 2, f"=AVERAGE({col})", "prix")
        ecrire(ws, r, 3, f'=AVERAGEIFS({col},{HC(C_PER)},"<=3")', "prix")
        ecrire(ws, r, 4, f'=AVERAGEIFS({col},{HC(C_PER)},">=4",{HC(C_PER)},"<=6")', "prix")
        ecrire(ws, r, 5, f"=AVERAGEIFS({col},{HC(C_PER)},7)", "prix")
        ecrire(ws, r, 6, f"=AVERAGEIFS({col},{HC(C_PER)},8)", "prix")
    ecrire(ws, 39, 1, "Repère : forward Base × (1 − prime)", font=F_SOUS)
    for j, prod in enumerate(["Cal-27", "Q1-27", "Q2-27", "Q3-27", "Q4-27"]):
        ecrire(ws, 39, 2 + j, f"={A(S_PAR, 2, ROW_PRIX[prod])}*(1-prime)", "prix")

    section(ws, 41, "Spot moyen de chaque jour (€/MWh)", 1, 7)
    entetes(ws, 42, 1, ["Date"] + [f"Météo {a}" for a in annees])
    jours = pd.Series(range(N), index=pd.DatetimeIndex(H.index.tz_localize(None)).normalize())
    grp = jours.groupby(level=0)
    D0 = 43
    for k, (jour, s) in enumerate(grp):
        r = D0 + k
        r1, r2 = H0 + int(s.min()), H0 + int(s.max())
        ecrire(ws, r, 1, jour.to_pydatetime(), "date")
        for i in range(NY):
            c = L(SC(i) + 8)
            ecrire(ws, r, 2 + i, f"=AVERAGE({q(S_H)}!{c}{r1}:{c}{r2})", "prix")
    DN = D0 + len(grp) - 1
    largeurs(ws, [18, 12, 12, 12, 12, 12, 16, 12, 12, 12, 12, 12])
    lc = LineChart()
    lc.title = "Spot 2027 : moyenne journalière par scénario météo"
    lc.y_axis.title = "€/MWh"
    lc.add_data(Reference(ws, min_col=2, max_col=1 + NY, min_row=42, max_row=DN), titles_from_data=True)
    lc.set_categories(Reference(ws, min_col=1, min_row=D0, max_row=DN))
    lc.x_axis.number_format = "mmm"
    lc.height, lc.width = 10, 22
    for s_ in lc.series:
        s_.smooth = False
        s_.graphicalProperties.line.width = 12000
    ws.add_chart(lc, "I33")
    NIV = lambda i: RG(S_SP, 9 + i, 16, 9 + i, 31)

    # ------------------------------------------------------------------
    # 1c PRIX T2 ANTICIPÉ (ALPHA)
    # ------------------------------------------------------------------
    ws = ws_al
    titre(ws, "Étape 1c : prix du 2e achat anticipé le 22/09 (méthode alpha)",
          "F_T2 = F_T1 + α × (niveau du scénario − F_T1). Contrôle CRE : 1 écart-type = F × 20 % × √(70/365).")
    section(ws, 3, "Prix Base anticipé au 01/12 pour chaque produit et chaque scénario", 1, 19)
    entetes(ws, 4, 1, ["Produit", "Prix 22/09"] + [f"Anomalie {a}" for a in annees] + [f"Niveau {a}" for a in annees]
            + [f"F_T2 anticipé {a}" for a in annees] + ["Min", "Max", "Écart max au prix du 22/09",
                                                         "1 écart-type CRE (70 jours)", "Dans ±1 écart-type ?"], 45)
    for k, prod in enumerate(produits):
        r = 5 + k
        ecrire(ws, r, 1, prod)
        if k >= 5:
            rq = ROW_PRIX[f"Q{(k - 5) // 3 + 1}-27"]
            cm, cq = A(S_PAR, 2, ROW_PRIX[prod]), A(S_PAR, 2, rq)
            ecrire(ws, r, 2, f"=IF(ISBLANK({cm}),{cq},{cm})", "prix")
        else:
            ecrire(ws, r, 2, "=" + A(S_PAR, 2, ROW_PRIX[prod]), "prix")
        for i in range(NY):
            if k == 0:
                an = A(S_CO, 4, 21 + i)
            elif k <= 4:
                an = A(S_CO, 7 + i, 49 + k)
            else:
                an = A(S_CO, 6 + i, 38 + (k - 4))
            ecrire(ws, r, 3 + i, "=" + an, "0.00%")
            ecrire(ws, r, 7 + i, f"=$B{r}*(1-prime)*(1+b_util*{L(3 + i)}{r})", "prix")
            ecrire(ws, r, 11 + i, f"=$B{r}+alpha*({L(7 + i)}{r}-$B{r})", "prix")
        ecrire(ws, r, 15, f"=MIN(K{r}:N{r})", "prix")
        ecrire(ws, r, 16, f"=MAX(K{r}:N{r})", "prix")
        ecrire(ws, r, 17, f"=MAX(ABS(O{r}-B{r}),ABS(P{r}-B{r}))", "prix")
        ecrire(ws, r, 18, f"=B{r}*volatilite*SQRT(jours_t1_t2/365)", "prix")
        ecrire(ws, r, 19, f'=IF(Q{r}<=R{r},"oui","non")')

    section(ws, 17, "Sensibilité du Cal anticipé à alpha", 1, 7)
    entetes(ws, 18, 1, ["Alpha"] + [f"Cal anticipé {a}" for a in annees])
    for j, aval in enumerate([0.2, "=alpha", 0.6]):
        r = 19 + j
        entree(ws, r, 1, aval, "0.00") if j != 1 else ecrire(ws, r, 1, aval, "0.00")
        for i in range(NY):
            ecrire(ws, r, 2 + i, f"=$B$5+$A{r}*({L(7 + i)}$5-$B$5)", "prix")

    section(ws, 23, "Pourquoi alpha ≈ 0,3 à 0,5 : part du temps écoulé entre le 22/09 et la livraison", 1, 8)
    entetes(ws, 24, 1, ["Trimestre", "Milieu de la livraison", "Jours depuis le 22/09", "α = 70 / jours",
                        "Énergie normale (MWh)", "Prix Base 22/09", "Poids = énergie × prix"], 45)
    milieux = [date(2027, 2, 15), date(2027, 5, 16), date(2027, 8, 15), date(2027, 11, 15)]
    for k in range(4):
        r = 25 + k
        ecrire(ws, r, 1, f"Q{k + 1}-27")
        entree(ws, r, 2, milieux[k], "date")
        ecrire(ws, r, 3, f"=B{r}-DATE(2026,9,22)", "int")
        ecrire(ws, r, 4, f"=jours_t1_t2/C{r}", "0.00")
        ecrire(ws, r, 5, "=" + A(S_CO, 2, 50 + k), "mwh")
        ecrire(ws, r, 6, "=" + A(S_PAR, 2, ROW_PRIX[f"Q{k + 1}-27"]), "prix")
        ecrire(ws, r, 7, f"=E{r}*F{r}", "eur")
    ecrire(ws, 30, 1, "α pondéré par le coût des achats")
    resultat(ws, 30, 2, "=SUMPRODUCT(D25:D28,G25:G28)/SUM(G25:G28)", "0.00")
    ecrire(ws, 31, 1, "Part de l'été (Q2 + Q3) dans le coût")
    ecrire(ws, 31, 2, "=(G26+G27)/SUM(G25:G28)", "0%")
    ecrire(ws, 32, 1, "α au-delà duquel S4 protège mieux que S7 (≈ 2 × part de l'été)")
    ecrire(ws, 32, 2, "=2*B31", "0.00")
    note(ws, 33, 1, "Si l'information arrive à rythme régulier, le forward a fait en moyenne la part "
                    "« temps écoulé / temps jusqu'au prix final » du chemin : c'est la même hypothèse que la règle √temps du cours.")
    largeurs(ws, [40, 14] + [12] * 17)

    # ------------------------------------------------------------------
    # 1d STRATÉGIES
    # ------------------------------------------------------------------
    ws = ws_st
    titre(ws, "Étape 1d : stratégies de couverture en % et en MW par produit",
          "Achat du 22/09 : un % de la couverture 100 % de chaque période. Achat du 01/12 : on complète jusqu'au % visé.")
    section(ws, 3, "Couverture 100 % par période (portefeuille du 22/09)", 1, 6)
    entetes(ws, 4, 1, ["Période", "Produit", "Base 100 % (MW)", "Peak 100 % (MW)", "Heures", "Heures peak"])
    for p in range(1, 9):
        r = 4 + p
        ecrire(ws, r, 1, NOMS_P[p - 1])
        ecrire(ws, r, 2, PROD_P[p - 1])
        ecrire(ws, r, 3, "=" + A(S_CO, 8, 27 + p), "mw")
        ecrire(ws, r, 4, "=" + A(S_CO, 9, 27 + p), "mw")
        ecrire(ws, r, 5, "=" + A(S_BL, 3, 23 + p), "int")
        ecrire(ws, r, 6, "=" + A(S_BL, 4, 23 + p), "int")

    section(ws, 14, "Définition des stratégies (modifiable)", 1, 18)
    ecrire(ws, 15, 2, "Achat du 22/09 : % de la couverture 100 %", bold=True)
    ecrire(ws, 15, 10, "Achat du 01/12", bold=True)
    entetes(ws, 16, 1, ["Stratégie"] + [f"T1 {x}" for x in PROD_P] + ["Compléter le 01/12 ? (1/0)"]
            + [f"Viser {x}" for x in PROD_P], 45)
    ST0 = 17
    for s_i, s in enumerate(strats):
        r = ST0 + s_i
        ecrire(ws, r, 1, s["nom"], bold=True)
        for p in range(8):
            entree(ws, r, 2 + p, s["r1"][p], "0%")
            entree(ws, r, 11 + p, s["r2"][p], "0%")
        entree(ws, r, 10, s["actif"], "0")
    NOMS_ST = RG(S_ST, 1, ST0, 1, ST0 + NS - 1)

    blocs_mw = [("MW Base achetés le 22/09", 24), ("MW Peak achetés le 22/09", 35),
                ("MW Base achetés le 01/12 (prévus)", 46), ("MW Peak achetés le 01/12 (prévus)", 57),
                ("MW Base total", 68), ("MW Peak total", 79)]
    for lib, r0 in blocs_mw:
        section(ws, r0, lib, 1, 1 + NS)
        entetes(ws, r0 + 1, 1, ["Produit"] + [f"={A(S_ST, 1, ST0 + s_i)}" for s_i in range(NS)])
        for p in range(1, 9):
            ecrire(ws, r0 + 1 + p, 1, PROD_P[p - 1])
    MWB1, MWP1, MWB2, MWP2, MWBT, MWPT = 26, 37, 48, 59, 70, 81     # première ligne de chaque matrice
    for s_i in range(NS):
        c = 2 + s_i
        rs = ST0 + s_i
        for p in range(1, 9):
            ratio1, cible = f"{q(S_ST)}!${L(1 + p)}${rs}", f"${L(10 + p)}${rs}"
            ecrire(ws, MWB1 + p - 1, c, f"={ratio1}*$C${4 + p}", "mw")
            ecrire(ws, MWP1 + p - 1, c, f"={ratio1}*$D${4 + p}", "mw")
            ecrire(ws, MWB2 + p - 1, c, f"=IF($J${rs}=1,{cible}*$C${4 + p}-{L(c)}{MWB1 + p - 1},0)", "mw")
            ecrire(ws, MWP2 + p - 1, c, f"=IF($J${rs}=1,{cible}*$D${4 + p}-{L(c)}{MWP1 + p - 1},0)", "mw")
            ecrire(ws, MWBT + p - 1, c, f"={L(c)}{MWB1 + p - 1}+{L(c)}{MWB2 + p - 1}", "mw")
            ecrire(ws, MWPT + p - 1, c, f"={L(c)}{MWP1 + p - 1}+{L(c)}{MWP2 + p - 1}", "mw")

    section(ws, 90, "Résumé par stratégie", 1, 12)
    entetes(ws, 91, 1, ["Stratégie", "Énergie achetée le 22/09 (MWh)", "Énergie prévue le 01/12 (MWh)",
                        "Couverture 22/09", "Couverture totale", "Couverture Q1", "Couverture Q2",
                        "Couverture Q3", "Couverture Q4", "Dépasse 100 % d'un trimestre (1/0)",
                        "Coût de l'achat du 22/09 (€)", "Prix moyen du 22/09 (€/MWh)"], 45)
    SUM0 = 92
    hp_loc, hpk_loc = "$E$5:$E$12", "$F$5:$F$12"
    for s_i in range(NS):
        r = SUM0 + s_i
        c = L(2 + s_i)
        ecrire(ws, r, 1, f"={L(1)}{ST0 + s_i}", bold=True)
        ecrire(ws, r, 2, f"=SUMPRODUCT({c}{MWB1}:{c}{MWB1 + 7},{hp_loc})+SUMPRODUCT({c}{MWP1}:{c}{MWP1 + 7},{hpk_loc})", "mwh")
        ecrire(ws, r, 3, f"=SUMPRODUCT({c}{MWB2}:{c}{MWB2 + 7},{hp_loc})+SUMPRODUCT({c}{MWP2}:{c}{MWP2 + 7},{hpk_loc})", "mwh")
        ecrire(ws, r, 4, f"=B{r}/{A(S_CO, 2, 17)}", "pct")
        ecrire(ws, r, 5, f"=(B{r}+C{r})/{A(S_CO, 2, 17)}", "pct")
        for qq in range(1, 5):
            ps = PERIODES_Q[qq]
            a_, b_ = ps[0] - 1, ps[-1] - 1
            ecrire(ws, r, 5 + qq, (f"=(SUMPRODUCT({c}{MWBT + a_}:{c}{MWBT + b_},$E${5 + a_}:$E${5 + b_})"
                                   f"+SUMPRODUCT({c}{MWPT + a_}:{c}{MWPT + b_},$F${5 + a_}:$F${5 + b_}))"
                                   f"/{A(S_CO, 2, 49 + qq)}"), "pct")
        ecrire(ws, r, 10, f"=IF(MAX(F{r}:I{r})>1.0005,1,0)", "int")
        ecrire(ws, r, 11, (f"=SUMPRODUCT({c}{MWB1}:{c}{MWB1 + 7},{hp_loc},{FB1})"
                           f"+SUMPRODUCT({c}{MWP1}:{c}{MWP1 + 7},{hpk_loc},{FP1})"), "eur")
        ecrire(ws, r, 12, f'=IF(B{r}>0,K{r}/B{r},"")', "prix")
    largeurs(ws, [24] + [12] * 18)

    # ------------------------------------------------------------------
    # 1e MARGES T1
    # ------------------------------------------------------------------
    ws = ws_m1
    titre(ws, "Étape 1e : marges de chaque stratégie dans chaque scénario, et règle du Board",
          "Marge = ventes − achat du 22/09 − achat du 01/12 (prix anticipé) − achats et reventes au spot.")
    section(ws, 4, "Données par bloc (16 blocs = 8 périodes × peak / hors peak)", 1, 4 + 2 * NY + NS)
    entetes(ws, 5, 1, ["Bloc", "Période", "Peak", "Heures"] + [f"Σ conso × forme {a} (MWh)" for a in annees]
            + [f"Niveau de prix {a}" for a in annees] + [f"MW livrés {s['nom'].split()[0]}" for s in strats], 45)
    for b in range(1, 17):
        r = 5 + b
        p, k = (b - 1) // 2 + 1, (b - 1) % 2
        ecrire(ws, r, 1, b)
        ecrire(ws, r, 2, p)
        ecrire(ws, r, 3, k)
        ecrire(ws, r, 4, "=" + A(S_BL, 5, 4 + b), "int")
        for i in range(NY):
            c0 = 8 + 6 * i
            ecrire(ws, r, 5 + i, f"={A(S_BL, c0 + 2, 4 + b)}+{A(S_BL, c0 + 3, 4 + b)}", "mwh")
            ecrire(ws, r, 5 + NY + i, "=" + A(S_SP, 9 + i, 15 + b), "prix")
        for s_i in range(NS):
            cs = L(2 + s_i)
            ecrire(ws, r, 5 + 2 * NY + s_i,
                   (f"=INDEX({q(S_ST)}!${cs}${MWBT}:${cs}${MWBT + 7},B{r})"
                    f"+C{r}*INDEX({q(S_ST)}!${cs}${MWPT}:${cs}${MWPT + 7},B{r})"), "mw")
    LS1 = lambda i: f"${L(5 + i)}$6:${L(5 + i)}$21"
    NV1 = lambda i: f"${L(5 + NY + i)}$6:${L(5 + NY + i)}$21"
    LV1 = lambda s_i: f"${L(5 + 2 * NY + s_i)}$6:${L(5 + 2 * NY + s_i)}$21"

    section(ws, 23, "Volume de chaque scénario (MWh)", 1, 4)
    for i, a in enumerate(annees):
        r = 24 + i
        c0 = 8 + 6 * i
        ecrire(ws, r, 1, f"Météo {a}")
        ecrire(ws, r, 2, f"=SUM({BL(c0)})+SUM({BL(c0 + 1)})", "mwh")

    G0 = 31
    section(ws, 29, "Marge de chaque stratégie dans chaque scénario", 1, 11)
    entetes(ws, 30, 1, ["Stratégie", "Prix", "Hausse du spot", "Scénario", "Volume (MWh)", "Ventes (€)",
                        "Achat 22/09 (€)", "Achat 01/12 (€)", "Spot : achats − reventes (€)", "Marge (€)",
                        "Marge (€/MWh)"], 45)
    chocs_lib = [("Météo seule", "0"), ("Tension", "=choc_tension"), ("Crise", "=choc_crise")]
    for s_i in range(NS):
        cs = L(2 + s_i)
        for c_i, (lib, val) in enumerate(chocs_lib):
            for i in range(NY):
                r = G0 + s_i * 12 + c_i * 4 + i
                ecrire(ws, r, 1, f"={A(S_ST, 1, ST0 + s_i)}")
                ecrire(ws, r, 2, lib)
                ecrire(ws, r, 3, float(val) if val == "0" else val, "0%")
                ecrire(ws, r, 4, annees[i])
                ecrire(ws, r, 5, f"=$B${24 + i}", "mwh")
                ecrire(ws, r, 6, f"=E{r}*prix_vente", "eur")
                ecrire(ws, r, 7, f"={A(S_ST, 11, SUM0 + s_i)}", "eur")
                an = ANOM_P(i)
                t2b = f"{q(S_ST)}!${cs}${MWB2}:${cs}${MWB2 + 7}"
                t2p = f"{q(S_ST)}!${cs}${MWP2}:${cs}${MWP2 + 7}"
                ecrire(ws, r, 8, (f"=SUMPRODUCT({t2b},{HP},{FB1}*(1-alpha)+alpha*{FB1}*(1-prime)*(1+$C{r})*(1+b_util*{an}))"
                                  f"+SUMPRODUCT({t2p},{HPK},{FP1}*(1-alpha)+alpha*{FP1}*(1-prime)*(1+$C{r})*(1+b_util*{an}))"), "eur")
                ecrire(ws, r, 9, (f"=(1+$C{r})*(SUMPRODUCT({NV1(i)},{LS1(i)})"
                                  f"-SUMPRODUCT({NV1(i)},{LV1(s_i)},$D$6:$D$21))"), "eur")
                ecrire(ws, r, 10, f"=F{r}-G{r}-H{r}-I{r}", "eur")
                ecrire(ws, r, 11, f"=J{r}/E{r}", "prix")
    GN = G0 + NS * 12 - 1

    def synthese(ws, r0, noms_ref, g0, dep_refs, titre_sec, nom_choix):
        section(ws, r0, titre_sec, 1, 11)
        entetes(ws, r0 + 1, 1, ["Option", "Marge espérée (€/MWh)", "Pire cas météo", "Pire cas tension",
                                "Pire cas crise", "Pire cas retenu", "Dépasse 100 % (1/0)",
                                "Tient le plancher (1/0)", "Pire cas si permis", "Candidate (1/0)", "Score"], 45)
        s0 = r0 + 2
        sN = s0 + NS - 1
        for s_i in range(NS):
            r = s0 + s_i
            b0 = g0 + s_i * 12
            ecrire(ws, r, 1, noms_ref[s_i], bold=True)
            ecrire(ws, r, 2, f"=SUMPRODUCT({PROBAS},K{b0}:K{b0 + 3})/SUM({PROBAS})", "prix")
            ecrire(ws, r, 3, f"=MIN(K{b0}:K{b0 + 3})", "prix")
            ecrire(ws, r, 4, f"=MIN(K{b0 + 4}:K{b0 + 7})", "prix")
            ecrire(ws, r, 5, f"=MIN(K{b0 + 8}:K{b0 + 11})", "prix")
            ecrire(ws, r, 6, f"=MIN(C{r},IF(compte_tension=1,D{r},1E9),IF(compte_crise=1,E{r},1E9))", "prix")
            ecrire(ws, r, 7, dep_refs[s_i], "int")
            ecrire(ws, r, 8, f"=IF(AND(F{r}>=plancher,G{r}=0),1,0)", "int")
            ecrire(ws, r, 9, f'=IF(G{r}=1,"",F{r})', "prix")
            ecrire(ws, r, 10, f"=IF($B${sN + 2}>0,H{r},IF(AND(G{r}=0,F{r}>=MAX($I${s0}:$I${sN})-0.01),1,0))", "int")
            ecrire(ws, r, 11, f'=IF(J{r}=1,B{r},"")', "prix")
        r = sN + 2
        ecrire(ws, r, 1, "Nombre d'options qui tiennent le plancher")
        ecrire(ws, r, 2, f"=SUM(H{s0}:H{sN})", "int")
        ecrire(ws, r + 1, 1, nom_choix, bold=True)
        resultat(ws, r + 1, 2, f"=INDEX(A{s0}:A{sN},MATCH(MAX(K{s0}:K{sN}),K{s0}:K{sN},0))")
        ecrire(ws, r + 2, 1, "Règle appliquée")
        ecrire(ws, r + 2, 2, f'=IF(B{r}>0,"Meilleure marge espérée parmi les options qui tiennent le plancher",'
                             f'"Aucune option ne tient le plancher : la plus protectrice (meilleur pire cas, à 0,01 près)")')
        ecrire(ws, r + 3, 1, "Marge espérée de l'option retenue (€/MWh)")
        resultat(ws, r + 3, 2, f"=INDEX(B{s0}:B{sN},MATCH(B{r + 1},A{s0}:A{sN},0))", "prix")
        ecrire(ws, r + 4, 1, "Pire cas de l'option retenue (€/MWh)")
        resultat(ws, r + 4, 2, f"=INDEX(F{s0}:F{sN},MATCH(B{r + 1},A{s0}:A{sN},0))", "prix")
        return s0, sN, r + 1

    SY1 = GN + 3
    s0_1, sN_1, CH1 = synthese(ws, SY1, [f"={A(S_ST, 1, ST0 + s_i)}" for s_i in range(NS)], G0,
                               [f"={A(S_ST, 10, SUM0 + s_i)}" for s_i in range(NS)],
                               "Règle du Board : marge espérée, pires cas et choix", "Stratégie retenue le 22/09")
    rp = CH1 + 6
    section(ws, rp, "Probabilité des chocs de prix (modèle log-normal, volatilité CRE)", 1, 6)
    ecrire(ws, rp + 1, 1, "Écart-type du niveau de prix, vu du 22/09")
    ecrire(ws, rp + 1, 2, "=volatilite*SQRT(horizon_t1)", "0.0%")
    ecrire(ws, rp + 2, 1, "Probabilité d'une hausse ≥ choc tension")
    ecrire(ws, rp + 2, 2, f"=1-NORMSDIST((LN(1+choc_tension)+B{rp + 1}^2/2)/B{rp + 1})", "0.0%")
    ecrire(ws, rp + 3, 1, "Probabilité d'une hausse ≥ choc crise")
    ecrire(ws, rp + 3, 2, f"=1-NORMSDIST((LN(1+choc_crise)+B{rp + 1}^2/2)/B{rp + 1})", "0.0%")
    note(ws, rp + 4, 1, "P = 1 − Φ((ln(1 + choc) + σ²/2) / σ). Ce modèle sous-estime les vraies crises (2021-2022).")
    largeurs(ws, [46, 14, 12, 12, 14, 14, 14, 14, 16, 14, 14, 13, 13, 13, 13, 13, 13])
    bc = BarChart()
    bc.type = "col"
    bc.title = "Marge espérée et pire cas retenu (€/MWh)"
    bc.add_data(Reference(ws, min_col=2, min_row=s0_1 - 1, max_row=sN_1), titles_from_data=True)
    bc.add_data(Reference(ws, min_col=6, min_row=s0_1 - 1, max_row=sN_1), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=s0_1, max_row=sN_1))
    bc.height, bc.width = 9, 18
    ws.add_chart(bc, f"M{SY1}")

    # ------------------------------------------------------------------
    # 2 DÉCISION T2
    # ------------------------------------------------------------------
    ws = ws_t2
    titre(ws, "Étape 2 : décision du 01/12 (achat du 22/09 figé, prix et clients réels du 01/12)",
          "On garde l'achat du 22/09 tel quel et on compare les façons de faire le 2e achat, avec la même règle.")
    section(ws, 3, "Mise à jour du portefeuille", 1, 4)
    entetes(ws, 4, 1, ["", "22/09", "01/12", ""])
    ecrire(ws, 5, 1, "Stratégie décidée le 22/09 (achat figé)")
    resultat(ws, 5, 2, f"={A(S_M1, 2, CH1)}")
    ecrire(ws, 6, 1, "Clients RES1")
    ecrire(ws, 6, 2, "=" + A(S_PAR, 2, 6), "#,##0")
    ecrire(ws, 6, 3, "=" + A(S_PAR, 3, 6), "#,##0")
    ecrire(ws, 7, 1, "Clients PRO1")
    ecrire(ws, 7, 2, "=" + A(S_PAR, 2, 7), "#,##0")
    ecrire(ws, 7, 3, "=" + A(S_PAR, 3, 7), "#,##0")
    ecrire(ws, 8, 1, "Volume normal 2027 (MWh)")
    ecrire(ws, 8, 2, "=" + A(S_CO, 2, 17), "mwh")
    ecrire(ws, 8, 3, "=SUM(B40:B47)", "mwh")
    ecrire(ws, 9, 1, "Variation du volume prévu")
    ecrire(ws, 9, 3, "=C8/B8-1", "0.0%")
    ecrire(ws, 10, 1, "Énergie achetée le 22/09 (MWh)")
    ecrire(ws, 10, 2, f"=INDEX({RG(S_ST, 2, SUM0, 2, SUM0 + NS - 1)},MATCH(B5,{RG(S_ST, 1, SUM0, 1, SUM0 + NS - 1)},0))", "mwh")
    ecrire(ws, 11, 1, "Taux de couverture après le 22/09")
    ecrire(ws, 11, 2, "=B10/B8", "pct")
    ecrire(ws, 11, 3, "=B10/C8", "pct")
    note(ws, 12, 1, "Le portefeuille a baissé : la même énergie achetée couvre une plus grande part du nouveau volume.")

    section(ws, 13, "Prix EEX du 01/12 par période", 1, 7)
    entetes(ws, 14, 1, ["Période", "Produit", "Forward Base", "Forward Peak", "Heures", "Heures peak", "Hors peak déduit"])
    for p in range(1, 9):
        r = 14 + p
        ecrire(ws, r, 1, NOMS_P[p - 1])
        ecrire(ws, r, 2, PROD_P[p - 1])
        ecrire(ws, r, 3, prix_periode(p, 4), "prix")
        ecrire(ws, r, 4, prix_periode(p, 5), "prix")
        ecrire(ws, r, 5, "=" + A(S_BL, 3, 23 + p), "int")
        ecrire(ws, r, 6, "=" + A(S_BL, 4, 23 + p), "int")
        ecrire(ws, r, 7, f"=(C{r}*E{r}-D{r}*F{r})/(E{r}-F{r})", "prix")
    FB2, FP2 = "$C$15:$C$22", "$D$15:$D$22"

    section(ws, 24, "Prix anticipé le 22/09 et prix réel du 01/12 (Base)", 1, 10)
    entetes(ws, 25, 1, ["Produit", "Prix 22/09", "Anticipé min", "Anticipé max", "Réel 01/12",
                        "Écart réel − 22/09 (€/MWh)", "Écart (%)", "Dans la fourchette anticipée ?",
                        "1 écart-type CRE", "Écart en écarts-types"], 45)
    for k, prod in enumerate(produits):
        r = 26 + k
        ecrire(ws, r, 1, prod)
        ecrire(ws, r, 2, "=" + A(S_AL, 2, 5 + k), "prix")
        ecrire(ws, r, 3, "=" + A(S_AL, 15, 5 + k), "prix")
        ecrire(ws, r, 4, "=" + A(S_AL, 16, 5 + k), "prix")
        ecrire(ws, r, 5, "=" + A(S_PAR, 4, ROW_PRIX[prod]), "prix")
        ecrire(ws, r, 6, f"=E{r}-B{r}", "prix")
        ecrire(ws, r, 7, f"=E{r}/B{r}-1", "0.0%")
        ecrire(ws, r, 8, f'=IF(AND(E{r}>=C{r},E{r}<=D{r}),"oui","non")')
        ecrire(ws, r, 9, "=" + A(S_AL, 18, 5 + k), "prix")
        ecrire(ws, r, 10, f"=F{r}/I{r}", "0.0")

    section(ws, 38, "Couverture 100 % et anomalies avec les clients du 01/12", 1, 18)
    entetes(ws, 39, 1, ["Période", "Énergie normale (MWh)", "Conso moyenne hors peak (MW)", "Conso moyenne peak (MW)",
                        "Base 100 % (MW)", "Peak 100 % (MW)"] + [f"Énergie météo {a}" for a in annees]
            + [f"Anomalie brute {a}" for a in annees] + [f"Anomalie recentrée {a}" for a in annees], 45)
    for p in range(1, 9):
        r = 39 + p
        ro, rp_ = bl_ro(p), bl_ro(p) + 1
        ecrire(ws, r, 1, NOMS_P[p - 1])
        e_off = f"(k_res_t2*{A(S_BL, 6, ro)}+k_pro_t2*{A(S_BL, 7, ro)})"
        e_pk = f"(k_res_t2*{A(S_BL, 6, rp_)}+k_pro_t2*{A(S_BL, 7, rp_)})"
        ecrire(ws, r, 2, f"={e_off}+{e_pk}", "mwh")
        ecrire(ws, r, 3, f"={e_off}/{A(S_BL, 5, ro)}", "mw")
        ecrire(ws, r, 4, f"={e_pk}/{A(S_BL, 5, rp_)}", "mw")
        ecrire(ws, r, 5, f"=IF(D{r}>=C{r},C{r},B{r}/{A(S_BL, 3, 23 + p)})", "mw")
        ecrire(ws, r, 6, f"=IF(D{r}>=C{r},D{r}-C{r},0)", "mw")
        for i in range(NY):
            c0 = 8 + 6 * i
            ecrire(ws, r, 7 + i, (f"=k_res_t2*({A(S_BL, c0, ro)}+{A(S_BL, c0, rp_)})"
                                  f"+k_pro_t2*({A(S_BL, c0 + 1, ro)}+{A(S_BL, c0 + 1, rp_)})"), "mwh")
            ecrire(ws, r, 11 + i, f"={L(7 + i)}{r}/$B{r}-1", "0.00%")
            ecrire(ws, r, 15 + i, f"={L(11 + i)}{r}-AVERAGE($K{r}:$N{r})*recentrage", "0.00%")
    for qq in range(1, 5):
        r = 49 + qq
        ps = PERIODES_Q[qq]
        ecrire(ws, r, 1, f"Énergie normale Q{qq} (MWh)")
        ecrire(ws, r, 2, f"=SUM(B{39 + ps[0]}:B{39 + ps[-1]})", "mwh")
    ANOM2 = lambda i: f"${L(15 + i)}$40:${L(15 + i)}$47"

    section(ws, 55, "Blocs : niveaux de prix du 01/12, conso × forme, MW livrés par chaque option", 1, 13 + NS)
    entetes(ws, 56, 1, ["Bloc", "Période", "Peak", "Heures", "Forward du bloc"] + [f"Niveau {a}" for a in annees]
            + [f"Σ conso × forme {a}" for a in annees] + [f"MW livrés option {s['nom'].split()[0]}" for s in strats], 45)
    B20 = 57
    for b in range(1, 17):
        r = B20 + b - 1
        p, k = (b - 1) // 2 + 1, (b - 1) % 2
        ecrire(ws, r, 1, b)
        ecrire(ws, r, 2, p)
        ecrire(ws, r, 3, k)
        ecrire(ws, r, 4, "=" + A(S_BL, 5, 4 + b), "int")
        ecrire(ws, r, 5, f"=IF(C{r}=1,INDEX($D$15:$D$22,B{r}),INDEX($G$15:$G$22,B{r}))", "prix")
        for i in range(NY):
            c0 = 8 + 6 * i
            ecrire(ws, r, 6 + i, f"=$E{r}*(1-prime)*(1+b_util*INDEX({ANOM2(i)},B{r}))", "prix")
            ecrire(ws, r, 10 + i, f"=k_res_t2*{A(S_BL, c0 + 2, 4 + b)}+k_pro_t2*{A(S_BL, c0 + 3, 4 + b)}", "mwh")
    NV2 = lambda i: f"${L(6 + i)}${B20}:${L(6 + i)}${B20 + 15}"
    LS2 = lambda i: f"${L(10 + i)}${B20}:${L(10 + i)}${B20 + 15}"
    LV2 = lambda o: f"${L(14 + o)}${B20}:${L(14 + o)}${B20 + 15}"
    HB2 = f"$D${B20}:$D${B20 + 15}"
    for i, a in enumerate(annees):
        r = 75 + i
        c0 = 8 + 6 * i
        ecrire(ws, r, 1, f"Volume météo {a} (MWh)")
        ecrire(ws, r, 2, f"=k_res_t2*SUM({BL(c0)})+k_pro_t2*SUM({BL(c0 + 1)})", "mwh")

    section(ws, 80, "Achat du 22/09 figé (stratégie décidée le 22/09)", 1, 6)
    entetes(ws, 81, 1, ["Période", "Produit", "MW Base", "MW Peak", "Prix Base 22/09", "Prix Peak 22/09"])
    col_dec = f"MATCH($B$5,{RG(S_ST, 2, MWB1 - 1, 1 + NS, MWB1 - 1)},0)"
    for p in range(1, 9):
        r = 81 + p
        ecrire(ws, r, 1, NOMS_P[p - 1])
        ecrire(ws, r, 2, PROD_P[p - 1])
        ecrire(ws, r, 3, f"=INDEX({RG(S_ST, 2, MWB1, 1 + NS, MWB1 + 7)},{p},{col_dec})", "mw")
        ecrire(ws, r, 4, f"=INDEX({RG(S_ST, 2, MWP1, 1 + NS, MWP1 + 7)},{p},{col_dec})", "mw")
        ecrire(ws, r, 5, "=" + A(S_SP, 3, 4 + p), "prix")
        ecrire(ws, r, 6, "=" + A(S_SP, 4, 4 + p), "prix")
    ecrire(ws, 90, 1, "Coût de l'achat du 22/09 (€)", bold=True)
    resultat(ws, 90, 2, f"=SUMPRODUCT(C82:C89,{HP},E82:E89)+SUMPRODUCT(D82:D89,{HPK},F82:F89)", "eur")

    opt_noms = [f'="Règle T2 de "&{A(S_ST, 1, ST0 + o)}' for o in range(NS)]
    mats = [("MW Base achetés le 01/12 par option", 93), ("MW Peak achetés le 01/12 par option", 104),
            ("MW Base total (22/09 + 01/12)", 115), ("MW Peak total (22/09 + 01/12)", 126)]
    for lib, r0 in mats:
        section(ws, r0, lib, 1, 1 + NS)
        entetes(ws, r0 + 1, 1, ["Produit"] + opt_noms, 45)
        for p in range(1, 9):
            ecrire(ws, r0 + 1 + p, 1, PROD_P[p - 1])
    OB2, OP2, OBT, OPT = 95, 106, 117, 128
    for o in range(NS):
        c = 2 + o
        rs = ST0 + o
        for p in range(1, 9):
            cible = f"{q(S_ST)}!${L(10 + p)}${rs}"
            actif = f"{q(S_ST)}!$J${rs}"
            ecrire(ws, OB2 + p - 1, c, f"=IF({actif}=1,{cible}*$E${39 + p}-$C${81 + p},0)", "mw")
            ecrire(ws, OP2 + p - 1, c, f"=IF({actif}=1,{cible}*$F${39 + p}-$D${81 + p},0)", "mw")
            ecrire(ws, OBT + p - 1, c, f"=$C${81 + p}+{L(c)}{OB2 + p - 1}", "mw")
            ecrire(ws, OPT + p - 1, c, f"=$D${81 + p}+{L(c)}{OP2 + p - 1}", "mw")
        for b in range(1, 17):
            r = B20 + b - 1
            cs = L(c)
            ecrire(ws, r, 14 + o, f"=INDEX(${cs}${OBT}:${cs}${OBT + 7},B{r})+C{r}*INDEX(${cs}${OPT}:${cs}${OPT + 7},B{r})", "mw")

    section(ws, 137, "Résumé des options", 1, 10)
    entetes(ws, 138, 1, ["Option", "Énergie achetée le 01/12 (MWh)", "Coût de l'achat du 01/12 (€)",
                         "Prix moyen 01/12 (€/MWh)", "Couverture totale (% du volume du 01/12)",
                         "Couverture Q1", "Couverture Q2", "Couverture Q3", "Couverture Q4", "Dépasse 100 % (1/0)"], 45)
    OS0 = 139
    hp2, hpk2 = "$E$15:$E$22", "$F$15:$F$22"
    for o in range(NS):
        r = OS0 + o
        cs = L(2 + o)
        ecrire(ws, r, 1, f"={cs}94", bold=True)
        ecrire(ws, r, 2, f"=SUMPRODUCT({cs}{OB2}:{cs}{OB2 + 7},{hp2})+SUMPRODUCT({cs}{OP2}:{cs}{OP2 + 7},{hpk2})", "mwh")
        ecrire(ws, r, 3, f"=SUMPRODUCT({cs}{OB2}:{cs}{OB2 + 7},{hp2},{FB2})+SUMPRODUCT({cs}{OP2}:{cs}{OP2 + 7},{hpk2},{FP2})", "eur")
        ecrire(ws, r, 4, f'=IF(ABS(B{r})>0.001,C{r}/B{r},"")', "prix")
        ecrire(ws, r, 5, f"=($B$10+B{r})/$C$8", "pct")
        for qq in range(1, 5):
            ps = PERIODES_Q[qq]
            a_, b_ = ps[0] - 1, ps[-1] - 1
            ecrire(ws, r, 5 + qq, (f"=(SUMPRODUCT({cs}{OBT + a_}:{cs}{OBT + b_},$E${15 + a_}:$E${15 + b_})"
                                   f"+SUMPRODUCT({cs}{OPT + a_}:{cs}{OPT + b_},$F${15 + a_}:$F${15 + b_}))/$B${49 + qq}"), "pct")
        ecrire(ws, r, 10, f"=IF(MAX(F{r}:I{r})>1.0005,1,0)", "int")

    G2 = 147
    section(ws, 145, "Marge de chaque option dans chaque scénario (spot ré-ancré sur les prix du 01/12)", 1, 11)
    entetes(ws, 146, 1, ["Option", "Prix", "Hausse du spot", "Scénario", "Volume (MWh)", "Ventes (€)",
                         "Achat 22/09 (€)", "Achat 01/12 (€)", "Spot : achats − reventes (€)", "Marge (€)",
                         "Marge (€/MWh)"], 45)
    for o in range(NS):
        for c_i, (lib, val) in enumerate(chocs_lib):
            for i in range(NY):
                r = G2 + o * 12 + c_i * 4 + i
                ecrire(ws, r, 1, f"=$A${OS0 + o}")
                ecrire(ws, r, 2, lib)
                ecrire(ws, r, 3, float(val) if val == "0" else val, "0%")
                ecrire(ws, r, 4, annees[i])
                ecrire(ws, r, 5, f"=$B${75 + i}", "mwh")
                ecrire(ws, r, 6, f"=E{r}*prix_vente", "eur")
                ecrire(ws, r, 7, "=$B$90", "eur")
                ecrire(ws, r, 8, f"=$C${OS0 + o}", "eur")
                ecrire(ws, r, 9, f"=(1+$C{r})*(SUMPRODUCT({NV2(i)},{LS2(i)})-SUMPRODUCT({NV2(i)},{LV2(o)},{HB2}))", "eur")
                ecrire(ws, r, 10, f"=F{r}-G{r}-H{r}-I{r}", "eur")
                ecrire(ws, r, 11, f"=J{r}/E{r}", "prix")
    G2N = G2 + NS * 12 - 1
    SY2 = G2N + 3
    s0_2, sN_2, CH2 = synthese(ws, SY2, [f"=$A${OS0 + o}" for o in range(NS)], G2,
                               [f"=$J${OS0 + o}" for o in range(NS)],
                               "Règle du Board au 01/12 (même règle qu'au 22/09)", "Option retenue le 01/12")

    rr = CH2 + 6
    section(ws, rr, "Robustesse : choc de prix qui ferait passer chaque option sous le plancher", 1, 9)
    entetes(ws, rr + 1, 1, ["Option"] + [f"Choc critique {a}" for a in annees]
            + ["Choc critique (le plus faible)", "Probabilité de passer sous le plancher"], 45)
    ecrire(ws, rr + 2 + NS, 1, "Écart-type du niveau de prix, vu du 01/12")
    sig = f"$B${rr + 2 + NS}"
    ecrire(ws, rr + 2 + NS, 2, "=volatilite*SQRT(horizon_t2)", "0.0%")
    for o in range(NS):
        r = rr + 2 + o
        ecrire(ws, r, 1, f"=$A${OS0 + o}", bold=True)
        for i in range(NY):
            m0 = f"$K${G2 + o * 12 + i}"
            m1 = f"$K${G2 + o * 12 + 8 + i}"
            pente = f"(({m1}-{m0})/choc_crise)"
            ecrire(ws, r, 2 + i, f"=IF({m0}<plancher,0,IF({pente}>=0,9.99,(plancher-{m0})/{pente}))", "0%")
        ecrire(ws, r, 2 + NY, f"=MIN(B{r}:{L(1 + NY)}{r})", "0%")
        cc = f"{L(2 + NY)}{r}"
        ecrire(ws, r, 3 + NY, f"=IF({cc}<=0,1,IF({cc}>=9.99,0,1-NORMSDIST((LN(1+{cc})+{sig}^2/2)/{sig})))", "0.0%")
    note(ws, rr + 3 + NS, 1, "La marge baisse en ligne droite quand le prix monte : on cherche la hausse qui la ramène au plancher. "
                             "9,99 = jamais.")

    rf = rr + 5 + NS
    section(ws, rf, "Achats finaux par produit (option retenue)", 1, 7)
    entetes(ws, rf + 1, 1, ["Produit", "MW Base 22/09", "MW Peak 22/09", "MW Base 01/12", "MW Peak 01/12",
                            "MW Base total", "MW Peak total"])
    col_opt = f"MATCH($B${CH2},$B$94:${L(1 + NS)}$94,0)"
    for p in range(1, 9):
        r = rf + 1 + p
        ecrire(ws, r, 1, PROD_P[p - 1])
        ecrire(ws, r, 2, f"=C{81 + p}", "mw")
        ecrire(ws, r, 3, f"=D{81 + p}", "mw")
        ecrire(ws, r, 4, f"=INDEX($B${OB2}:${L(1 + NS)}${OB2 + 7},{p},{col_opt})", "mw")
        ecrire(ws, r, 5, f"=INDEX($B${OP2}:${L(1 + NS)}${OP2 + 7},{p},{col_opt})", "mw")
        ecrire(ws, r, 6, f"=B{r}+D{r}", "mw")
        ecrire(ws, r, 7, f"=C{r}+E{r}", "mw")
    RF0 = rf + 2
    ecrire(ws, rf + 10, 1, "Couverture totale atteinte (% du volume du 01/12)", bold=True)
    resultat(ws, rf + 10, 2, f"=INDEX($E${OS0}:$E${OS0 + NS - 1},MATCH($B${CH2},$A${OS0}:$A${OS0 + NS - 1},0))", "pct")
    largeurs(ws, [46, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 13, 13, 13, 13, 13, 13, 13])

    # ------------------------------------------------------------------
    # 3 POSITION OUVERTE 01/01/2027
    # ------------------------------------------------------------------
    ws = ws_t3
    titre(ws, "Étape 3 : position ouverte du 01/01/2027 et ordre day-ahead",
          "Position ouverte = consommation prévue − livraison des produits achetés − production solaire.")
    ecrire(ws, 4, 1, "Option retenue le 01/12")
    resultat(ws, 4, 2, f"={A(S_T2, 2, CH2)}")
    ecrire(ws, 5, 1, "MW Base livrés en janvier (22/09 + 01/12)")
    resultat(ws, 5, 2, f"={A(S_T2, 6, RF0)}", "mw")
    ecrire(ws, 6, 1, "MW Peak livrés en janvier (8h-20h, jours fériés compris)")
    resultat(ws, 6, 2, f"={A(S_T2, 7, RF0)}", "mw")
    note(ws, 7, 1, "Le 01/01/2027 est un vendredi férié : le Peak livre quand même, alors que nos clients consomment comme un dimanche.")
    entetes(ws, 9, 1, ["Heure", "Peak", "Coef RES1", "Coef PRO1", "T normale (°C)", "Gradient RES1", "Gradient PRO1",
                       "T prévue le 31/12 (°C)", "CM RES1", "CM PRO1", "Consommation prévue (MW)",
                       "Livraison couverte (MW)", "Solaire prévu (MW)", "Position ouverte (MW)", "Court ou long"], 45)
    t3 = d.get("t3", {})
    temp = list(t3.get("temp", [])) or [None] * 24
    sol = list(t3.get("solaire", [0] * 24)) if t3.get("solaireOn", True) else [0] * 24
    for h in range(24):
        r = 10 + h
        hr = H0 + h
        ecrire(ws, r, 1, f"{h:02d}:00 - {h + 1:02d}:00")
        ecrire(ws, r, 2, f"={q(S_H)}!D{hr}")
        ecrire(ws, r, 3, f"={q(S_H)}!F{hr}", "coef")
        ecrire(ws, r, 4, f"={q(S_H)}!G{hr}", "coef")
        ecrire(ws, r, 5, f"={q(S_H)}!K{hr}", "temp")
        ecrire(ws, r, 6, f"={q(S_H)}!L{hr}", "0.00000")
        ecrire(ws, r, 7, f"={q(S_H)}!M{hr}", "0.00000")
        entree(ws, r, 8, temp[h], "temp")
        for cc, g in ((9, "F"), (10, "G")):
            ecrire(ws, r, cc, (f"=IF(AND(H{r}<seuil,E{r}<seuil),1+{g}{r}*(E{r}-H{r}),"
                               f"IF(AND(H{r}<seuil,E{r}>=seuil),1+{g}{r}*(seuil-H{r}),"
                               f"IF(AND(E{r}<seuil,H{r}>=seuil),1+{g}{r}*(E{r}-seuil),1)))"), "coef")
        ecrire(ws, r, 11, f"=C{r}*pmoy_res_t3*I{r}+D{r}*pmoy_pro_t3*J{r}", "mw")
        ecrire(ws, r, 12, f"=$B$5+B{r}*$B$6", "mw")
        entree(ws, r, 13, sol[h], "mw")
        ecrire(ws, r, 14, f"=K{r}-L{r}-M{r}", "mw")
        ecrire(ws, r, 15, f'=IF(N{r}>0,"Court (achat)","Long (revente)")')
    ecrire(ws, 34, 1, "Total (MWh)", bold=True)
    for c in (11, 12, 13, 14):
        resultat(ws, 34, c, f"=SUM({L(c)}10:{L(c)}33)", "mw")
    ecrire(ws, 35, 1, "À acheter (heures courtes, MWh)")
    resultat(ws, 35, 14, '=SUMIF(N10:N33,">0")', "mw")
    ecrire(ws, 36, 1, "À revendre (heures longues, MWh)")
    resultat(ws, 36, 14, '=-SUMIF(N10:N33,"<0")', "mw")

    section(ws, 38, "Ordre day-ahead EPEX (format du bon de commande)", 1, 7)
    entetes(ws, 39, 1, ["Heure", "Net Load (MWh) = conso − livraison", "Production solaire (MWh)",
                        "Intervalle 1 : volume (prix < 0 : solaire arrêté)", "Intervalle 1 : prix limite (€/MWh)",
                        "Intervalle 2 : volume (Net Load − Production)", "Intervalle 2 : prix limite (€/MWh)"], 58)
    for h in range(24):
        r = 40 + h
        ecrire(ws, r, 1, f"=A{10 + h}")
        ecrire(ws, r, 2, f"=K{10 + h}-L{10 + h}", "mw")
        ecrire(ws, r, 3, f"=M{10 + h}", "mw")
        ecrire(ws, r, 4, f"=B{r}", "mw")
        ecrire(ws, r, 5, 0)
        ecrire(ws, r, 6, f"=B{r}-C{r}", "mw")
        ecrire(ws, r, 7, 4000)
    note(ws, 64, 1, "Achat = volume positif, revente = volume négatif. Coût marginal du solaire ≈ 0 €/MWh : "
                    "si le prix est négatif, on arrête la centrale et on achète tout le Net Load.")
    largeurs(ws, [44, 13, 12, 13, 13, 13, 13, 12, 9, 9, 13, 13, 12, 13, 15])
    oc = BarChart()
    oc.type = "col"
    oc.title = "Position ouverte du 01/01/2027 (MW) : > 0 on achète, < 0 on revend"
    oc.add_data(Reference(ws, min_col=14, min_row=9, max_row=33), titles_from_data=True)
    oc.set_categories(Reference(ws, min_col=1, min_row=10, max_row=33))
    oc.legend = None
    oc.height, oc.width = 9, 20
    ws.add_chart(oc, "Q9")


    # ------------------------------------------------------------------
    # 4 RÉSULTAT DU DAY-AHEAD DU 01/01/2027
    # ------------------------------------------------------------------
    ws = ws_r4
    da = getattr(config, "DA_0101", None)
    R4 = None                       # lignes clés, reprises dans le Récap
    titre(ws, "Étape 4 : résultat du day-ahead du 01/01/2027",
          "Prix publiés par le professeur, exécution de notre ordre (feuille 3), coût du day-ahead et marge de la journée.")
    if not (da and len(da.get("prix", [])) == 24):
        note(ws, 4, 1, "Résultats non disponibles : ajouter le bloc DA_0101 (prix du 01/01/2027) dans config.py, "
                       "puis relancer  python export_excel.py.")
        largeurs(ws, [110])
    else:
        fixe = list(da.get("fixe_par", [""] * 24)) + [""] * 24
        achats_g = list(da.get("achats_equipes", [None] * 24)) + [None] * 24
        T0 = 6                       # première heure du tableau horaire
        TT = T0 + 24                 # ligne des totaux
        section(ws, 4, f"Prix du day-ahead (groupe {da.get('groupe', '')}) et exécution de notre ordre", 1, 18)
        entetes(ws, 5, 1, ["Heure", "Prix day-ahead (€/MWh)", "Prix fixé par", "Achats nets du groupe (MWh)", "Peak (1/0)",
                           "Volume de notre ordre à ce prix (MWh)", "Part exécutée", "Volume exécuté (MWh)",
                           "Achat (MWh)", "Vente (MWh)", "Coût des achats (€)", "Recette des ventes (€)",
                           "Coût net pour nous (€)", "Consommation prévue (MW)", "Solaire (MW)",
                           "Livraison des produits (MW)", "Coût si tout acheté au spot (€)",
                           "Valeur au spot des produits livrés (€)"], 58)
        for h in range(24):
            r, r3, ro = T0 + h, 10 + h, 40 + h
            ecrire(ws, r, 1, f"={q(S_T3)}!A{r3}")
            donnee(ws, r, 2, float(da["prix"][h]), "prix")
            donnee(ws, r, 3, fixe[h])
            donnee(ws, r, 4, achats_g[h] if achats_g[h] is not None else "-", "0.0")
            ecrire(ws, r, 5, f"={q(S_T3)}!B{r3}", "int")
            # L'ordre a deux intervalles : volume 1 si prix < limite 1, sinon volume 2 si prix <= limite 2
            ecrire(ws, r, 6, f"=IF(B{r}<{q(S_T3)}!E{ro},{q(S_T3)}!D{ro},IF(B{r}<={q(S_T3)}!G{ro},{q(S_T3)}!F{ro},0))", "mw")
            entree(ws, r, 7, 1, "0%")
            ecrire(ws, r, 8, f"=F{r}*G{r}", "mw")
            ecrire(ws, r, 9, f"=MAX(H{r},0)", "mw")
            ecrire(ws, r, 10, f"=MAX(-H{r},0)", "mw")
            ecrire(ws, r, 11, f"=I{r}*B{r}", "eur")
            ecrire(ws, r, 12, f"=J{r}*B{r}", "eur")
            ecrire(ws, r, 13, f"=K{r}-L{r}", "eur")
            ecrire(ws, r, 14, f"={q(S_T3)}!K{r3}", "mw")
            ecrire(ws, r, 15, f"={q(S_T3)}!M{r3}", "mw")
            ecrire(ws, r, 16, f"={q(S_T3)}!L{r3}", "mw")
            ecrire(ws, r, 17, f"=(N{r}-O{r})*B{r}", "eur")
            ecrire(ws, r, 18, f"=P{r}*B{r}", "eur")
        ecrire(ws, TT, 1, "Total de la journée", bold=True)
        for c in (6, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18):
            resultat(ws, TT, c, f"=SUM({L(c)}{T0}:{L(c)}{TT - 1})", "eur" if c in (11, 12, 13, 17, 18) else "mw")
        ecrire(ws, TT + 1, 1, "Prix moyen payé à l'achat (€/MWh)")
        resultat(ws, TT + 1, 11, f"=IF(I{TT}>0,K{TT}/I{TT},0)", "prix")
        ecrire(ws, TT + 2, 1, "Prix moyen obtenu à la vente (€/MWh, négatif = on paie pour vendre)")
        resultat(ws, TT + 2, 12, f"=IF(J{TT}>0,L{TT}/J{TT},0)", "prix")
        note(ws, TT + 3, 1, "Part exécutée : 100 % par défaut. Si le professeur confirme qu'une partie de nos ventes n'a pas été "
                            "exécutée (heures à −600 €/MWh, ordres ajustés), mettre la part réelle dans la colonne G.")

        PRX, PK = f"$B${T0}:$B${TT - 1}", f"$E${T0}:$E${TT - 1}"
        ri = TT + 5
        section(ws, ri, "Indices de prix du jour", 1, 4)
        entetes(ws, ri + 1, 1, ["Indice", "€/MWh"])
        ind = [("Base (moyenne des 24 heures)", f"=AVERAGE({PRX})"),
               ("Peak (8h-20h)", f"=IF(SUM({PK})>0,SUMPRODUCT({PRX},{PK})/SUM({PK}),0)"),
               ("Hors peak (autres heures)", f"=IF(SUM({PK})<24,(SUM({PRX})-SUMPRODUCT({PRX},{PK}))/(24-SUM({PK})),0)"),
               ("Minimum", f"=MIN({PRX})"), ("Maximum", f"=MAX({PRX})"),
               ("Jan-27 Base coté la veille (31/12)", "=" + A(S_PAR, 6, ROW_PRIX["Jan-27"])),
               ("Jan-27 Peak coté la veille (31/12)", "=" + A(S_PAR, 7, ROW_PRIX["Jan-27"]))]
        for k, (lib, f) in enumerate(ind):
            ecrire(ws, ri + 2 + k, 1, lib)
            resultat(ws, ri + 2 + k, 2, f, "prix")

        rm = ri + 2 + len(ind) + 1
        section(ws, rm, "Offre de production du groupe (merit order) et coût par moyen marginal", 1, 9)
        entetes(ws, rm + 1, 1, ["Moyen", "Coût marginal (€/MWh)", "Capacité (MW)", "Capacité cumulée (MW)",
                                "Heures où il fixe le prix", "Nos achats (MWh)", "Nos ventes (MWh)",
                                "Coût net pour nous (€)", "Part du coût net"], 45)
        merit = list(da.get("merit", []))
        moyens = [(m[0], m[1], m[2]) for m in merit] + [("Ordre de vente d'une équipe à sa limite", da.get("plancher", -600), None)]
        FX = f"$C${T0}:$C${TT - 1}"
        for k, (nom, cm, cap) in enumerate(moyens):
            r = rm + 2 + k
            ecrire(ws, r, 1, nom)
            donnee(ws, r, 2, cm, "0")
            if cap is not None:
                donnee(ws, r, 3, cap, "0")
                ecrire(ws, r, 4, f"=SUM($C${rm + 2}:C{r})", "0")
            ecrire(ws, r, 5, f"=COUNTIF({FX},A{r})", "int")
            ecrire(ws, r, 6, f"=SUMIF({FX},A{r},$I${T0}:$I${TT - 1})", "mw")
            ecrire(ws, r, 7, f"=SUMIF({FX},A{r},$J${T0}:$J${TT - 1})", "mw")
            ecrire(ws, r, 8, f"=SUMIF({FX},A{r},$M${T0}:$M${TT - 1})", "eur")
            ecrire(ws, r, 9, f"=IF($M${TT}<>0,H{r}/$M${TT},0)", "pct")
        rb = rm + 2 + len(moyens)
        ecrire(ws, rb, 1, "Prix plancher / plafond du marché (€/MWh)")
        donnee(ws, rb, 2, da.get("plancher", -600), "0")
        donnee(ws, rb, 3, da.get("plafond", 4000), "#,##0")

        rj = rb + 2
        section(ws, rj, "Marge de la journée", 1, 5)
        entetes(ws, rj + 1, 1, ["Poste", "Volume (MWh)", "Prix (€/MWh)", "Montant (€)", "€/MWh consommé"])
        jan = rf + 2                                   # ligne « Jan-27 » des achats finaux (feuille 2)
        heures_pk = f"SUM({PK})"
        postes = [
            ("Ventes aux clients (consommation prévue)", f"=N{TT}", "=prix_vente", 1),
            ("Base achetée le 22/09 (Jan-27)", f"={A(S_T2, 2, jan)}*24", prix_periode(1, 2), -1),
            ("Peak acheté le 22/09 (Jan-27, heures peak du jour)", f"={A(S_T2, 3, jan)}*{heures_pk}", prix_periode(1, 3), -1),
            ("Base achetée le 01/12 (Jan-27)", f"={A(S_T2, 4, jan)}*24", prix_periode(1, 4), -1),
            ("Peak acheté le 01/12 (Jan-27, heures peak du jour)", f"={A(S_T2, 5, jan)}*{heures_pk}", prix_periode(1, 5), -1),
        ]
        r0j = rj + 2
        for k, (lib, v, pz, sg) in enumerate(postes):
            r = r0j + k
            ecrire(ws, r, 1, lib)
            ecrire(ws, r, 2, v, "mw")
            ecrire(ws, r, 3, pz, "prix")
            ecrire(ws, r, 4, f"={'' if sg > 0 else '-'}B{r}*C{r}", "eur")
            ecrire(ws, r, 5, f"=IF($B${r0j}>0,D{r}/$B${r0j},0)", "prix")
        rd = r0j + len(postes)
        ecrire(ws, rd, 1, "Day-ahead : achats − ventes (feuille ci-dessus)")
        ecrire(ws, rd, 2, f"=H{TT}", "mw")
        ecrire(ws, rd, 3, "—")
        ecrire(ws, rd, 4, f"=-M{TT}", "eur")
        ecrire(ws, rd, 5, f"=IF($B${r0j}>0,D{rd}/$B${r0j},0)", "prix")
        rmj = rd + 1
        ecrire(ws, rmj, 1, "Marge de la journée", bold=True)
        resultat(ws, rmj, 4, f"=SUM(D{r0j}:D{rd})", "eur")
        resultat(ws, rmj, 5, f"=IF($B${r0j}>0,D{rmj}/$B${r0j},0)", "prix")
        note(ws, rmj + 1, 1, "Hypothèse : consommation réelle = consommation prévue le 31/12 (la vraie consommation du 01/01 n'est pas connue).")

        rc = rmj + 3
        section(ws, rc, "Comparaisons", 1, 5)
        comp = [
            ("Coût des produits à terme livrés le 01/01 (€)", f"=-SUM(D{r0j + 1}:D{rd - 1})", "eur"),
            ("Valeur au spot de cette même énergie (€)", f"=R{TT}", "eur"),
            ("Prix moyen payé pour l'énergie livrée par les produits (€/MWh)", f"=IF(P{TT}>0,B{rc + 1}/P{TT},0)", "prix"),
            ("Coût si tout avait été acheté au day-ahead ce jour-là (€)", f"=Q{TT}", "eur"),
            ("Marge si tout avait été acheté au day-ahead (€)", f"=D{r0j}-Q{TT}", "eur"),
            ("Marge si tout avait été acheté au day-ahead (€/MWh)", f"=IF($B${r0j}>0,B{rc + 5}/$B${r0j},0)", "prix"),
            ("Solde net échangé au day-ahead (MWh)", f"=H{TT}", "mw"),
            ("Coût net du day-ahead (€)", f"=M{TT}", "eur"),
        ]
        for k, (lib, f, fm) in enumerate(comp):
            ecrire(ws, rc + 1 + k, 1, lib)
            resultat(ws, rc + 1 + k, 2, f, fm)
        note(ws, rc + 2 + len(comp), 1, "Une couverture se juge sur l'année : ce jour férié a été très bon marché, alors que "
                                        "l'énergie de janvier a été achetée au prix attendu pour tout le mois.")
        R4 = {"cout_da": f"={A(S_R4, 13, TT)}", "marge": f"={A(S_R4, 4, rmj)}", "marge_mwh": f"={A(S_R4, 5, rmj)}",
              "marge_spot": f"={A(S_R4, 2, rc + 5)}", "net": f"={A(S_R4, 8, TT)}"}

        largeurs(ws, [52, 14, 26, 13, 9, 14, 10, 13, 11, 11, 13, 13, 13, 13, 10, 13, 14, 14])
        ws.freeze_panes = "B6"
        for col, titre_g, pos in ((2, "Prix day-ahead du 01/01/2027 (€/MWh)", "T4"),
                                  (8, "Notre volume exécuté (MWh) : > 0 achat, < 0 vente", "T24")):
            g = BarChart()
            g.type = "col"
            g.title = titre_g
            g.add_data(Reference(ws, min_col=col, min_row=5, max_row=TT - 1), titles_from_data=True)
            g.set_categories(Reference(ws, min_col=1, min_row=T0, max_row=TT - 1))
            g.legend = None
            g.height, g.width = 9, 20
            ws.add_chart(g, pos)

    # ------------------------------------------------------------------
    # RÉCAP
    # ------------------------------------------------------------------
    ws = ws_rec
    titre(ws, "Récapitulatif de la stratégie 2027", "Toutes les valeurs viennent des feuilles de calcul.")
    section(ws, 3, "22/09 : première décision", 1, 4)
    rec = [
        (4, "Stratégie retenue", f"={A(S_M1, 2, CH1)}", None),
        (5, "Règle appliquée", f"={A(S_M1, 2, CH1 + 1)}", None),
        (6, "Marge espérée (€/MWh)", f"={A(S_M1, 2, CH1 + 2)}", "prix"),
        (7, "Pire cas (€/MWh)", f"={A(S_M1, 2, CH1 + 3)}", "prix"),
        (8, "Part achetée le 22/09 (% du volume normal)",
         f"=INDEX({RG(S_ST, 4, SUM0, 4, SUM0 + NS - 1)},MATCH(B4,{RG(S_ST, 1, SUM0, 1, SUM0 + NS - 1)},0))", "pct"),
    ]
    section(ws, 10, "01/12 : second achat", 1, 4)
    rec += [
        (11, "Option retenue", f"={A(S_T2, 2, CH2)}", None),
        (12, "Règle appliquée", f"={A(S_T2, 2, CH2 + 1)}", None),
        (13, "Marge espérée (€/MWh)", f"={A(S_T2, 2, CH2 + 2)}", "prix"),
        (14, "Pire cas (€/MWh)", f"={A(S_T2, 2, CH2 + 3)}", "prix"),
        (15, "Couverture totale (% du volume du 01/12)", f"={A(S_T2, 2, rf + 10)}", "pct"),
    ]
    section(ws, 17, "Hypothèses clés", 1, 4)
    rec += [
        (18, "Pente prix-volume b utilisée", "=b_util", "0.000"),
        (19, "R² de la régression", f"={A(S_B, 2, 8)}", "0.00"),
        (20, "b corrigé (test de sensibilité)", f"={A(S_B, 5, 6)}", "0.00"),
        (21, "Alpha", "=alpha", "0.00"),
        (22, "Alpha « temps écoulé » pondéré", f"={A(S_AL, 2, 30)}", "0.00"),
        (23, "Prime de risque", "=prime", "0.0%"),
        (24, "Plancher du Board (€/MWh)", "=plancher", "0.00"),
    ]
    section(ws, 26, "01/01/2027 : position ouverte", 1, 4)
    rec += [
        (27, "À acheter au day-ahead (MWh)", f"={A(S_T3, 14, 35)}", "mw"),
        (28, "À revendre au day-ahead (MWh)", f"={A(S_T3, 14, 36)}", "mw"),
        (29, "Position nette (MWh)", f"={A(S_T3, 14, 34)}", "mw"),
    ]
    for r, lib, f, fm in rec:
        ecrire(ws, r, 1, lib)
        resultat(ws, r, 2, f, fm)
    section(ws, 31, "MW achetés par produit", 1, 7)
    entetes(ws, 32, 1, ["Produit", "Base 22/09", "Peak 22/09", "Base 01/12", "Peak 01/12", "Base total", "Peak total"])
    for p in range(1, 9):
        r = 32 + p
        ecrire(ws, r, 1, PROD_P[p - 1])
        for j in range(6):
            ecrire(ws, r, 2 + j, f"={A(S_T2, 2 + j, rf + 1 + p)}", "mw")
    if R4:
        section(ws, 42, "01/01/2027 : résultat du day-ahead", 1, 4)
        for k, (lib, cle, fm) in enumerate([
                ("Solde net échangé au day-ahead (MWh)", "net", "mw"),
                ("Coût net du day-ahead (€)", "cout_da", "eur"),
                ("Marge de la journée (€)", "marge", "eur"),
                ("Marge de la journée (€/MWh)", "marge_mwh", "prix"),
                ("Marge si tout avait été acheté au day-ahead (€)", "marge_spot", "eur")]):
            ecrire(ws, 43 + k, 1, lib)
            resultat(ws, 43 + k, 2, R4[cle], fm)
    for r in (5, 12):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 28
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center")
    largeurs(ws, [44, 16, 12, 12, 12, 12, 12])

    # ------------------------------------------------------------------
    # LISEZ-MOI
    # ------------------------------------------------------------------
    ws = ws_lis
    titre(ws, "Market Game 2027 : le modèle en Excel",
          f"Généré par export_excel.py le {pd.Timestamp.now():%d/%m/%Y %H:%M}, à partir des mêmes données que le tableau de bord.")
    lignes = [
        ("Ce classeur", "Il refait, avec des formules Excel visibles, les étapes du modèle Python et du tableau de bord."),
        ("", "Changer une cellule bleue (feuille Paramètres ou Stratégies) recalcule tout le classeur."),
        ("Code couleur", "Bleu sur fond jaune : hypothèse modifiable. Fond bleu clair : donnée préparée par Python. "
                         "Fond vert : résultat clé. Blanc : formule."),
        ("", ""),
        ("Feuille", "Ce qu'elle fait"),
        (S_REC, "Les décisions et les chiffres clés."),
        (S_PAR, "Clients, θ, kVA, prix de vente, plancher, prime, alpha, b, chocs, probabilités, prix EEX."),
        (S_CO, "Étape 0 : volumes mensuels, partage peak / hors peak, couverture 100 % par période, anomalies."),
        (S_B, "Étape 1a : anomalies journalières de volume et de prix, pente b = Cov(x, y) / Var(x), R²."),
        (S_SP, "Étape 1b : prix par période, hors peak déduit, niveau de chaque bloc, courbes spot 2027."),
        (S_AL, "Étape 1c : prix du 2e achat anticipé (alpha), contrôle CRE, justification d'alpha."),
        (S_ST, "Étape 1d : stratégies en % et en MW par produit (Base, Peak), couverture, coût du 22/09."),
        (S_M1, "Étape 1e : marges (stratégie × scénario × choc), marge espérée, pires cas, règle du Board."),
        (S_T2, "Étape 2 : achat du 22/09 figé, prix et clients du 01/12, options du 2e achat, même règle."),
        (S_T3, "Étape 3 : position ouverte du 01/01/2027 heure par heure et ordre day-ahead EPEX."),
        (S_R4, "Étape 4 : prix day-ahead du 01/01/2027, exécution de notre ordre, coût du day-ahead et marge de la journée."),
        (S_H, "8 760 heures : coefficient de profil, CM météo, conso par scénario, forme et spot."),
        (S_BL, "Feuille technique : sommes par bloc (période × peak / hors peak)."),
        ("", ""),
        ("Préparé par Python", "Profils Enedis dé-météorisés (moyenne 2023-2025) et alignés sur 2027 par jour de semaine "
                               "et jours fériés ; écarts de température et prix historiques des années 2022 à 2025 "
                               "replacés sur 2027 ; données journalières de la régression."),
        ("Écart avec le tableau de bord", "Ici le CM est calculé heure par heure (tableau de bord : demi-heure par demi-heure). "
                                          "Les marges peuvent différer de quelques centimes d'€/MWh."),
        ("Non repris", "Le Monte Carlo reste dans le tableau de bord ; ici, la probabilité des chocs est calculée par formule."),
        ("Mise à jour", "Relancer  python export_excel.py  après un changement de données ou de config.py."),
    ]
    for k, (a_, b_) in enumerate(lignes):
        r = 4 + k
        c1 = ws.cell(row=r, column=1, value=a_)
        c2 = ws.cell(row=r, column=2, value=b_)
        c1.font = F_ENTETE if a_ else F_NORMAL
        c2.font = F_NORMAL
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        if a_ == "Feuille":
            c1.fill = c2.fill = FILL_ENTETE
    largeurs(ws, [30, 120])

    # Mise en page : paysage, une page en largeur ; zones d'impression limitées
    # pour les feuilles très longues (sinon des milliers de pages).
    zones = {S_H: "A1:AX60", S_B: f"A1:V{min(AN, 70)}", S_SP: "A1:T75", S_CO: "A1:T53"}
    from openpyxl.worksheet.properties import PageSetupProperties
    for w in wb.worksheets:
        w.sheet_view.zoomScale = 90
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        if w.title in zones:
            w.print_area = zones[w.title]
    return wb


def main():
    import os
    os.chdir(DOSSIER)          # le cache data/ est cherché à côté de ce script
    D, H, J, feries, annees = preparer_donnees()
    print("\nÉcriture du classeur Excel (cela peut prendre une minute)...")
    wb = construire_classeur(D, H, J, feries, annees)
    wb.save(SORTIE)
    print(f"Classeur écrit : {SORTIE}")
    print("À l'ouverture, Excel calcule toutes les formules (quelques secondes).")


if __name__ == "__main__":
    main()
